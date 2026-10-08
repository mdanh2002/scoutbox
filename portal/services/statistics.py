from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timedelta
from io import BytesIO
import re

from django.db.models import Q, Sum
from django.utils import timezone

from portal.models import Application, MailEvent, Opportunity, SearchProviderStat, UsageMetric


METRICS = [
    ('pages_scraped', 'Pages scraped', 'count', 'Discovery'),
    ('roles_found', 'Roles found', 'count', 'Discovery'),
    ('duplicates', 'Duplicates', 'count', 'Discovery'),
    ('already_applied', 'Already applied', 'count', 'Discovery'),
    ('high_priority', 'High priority', 'count', 'Funnel'),
    ('prepared', 'Applications prepared', 'count', 'Funnel'),
    ('newly_applied', 'Newly applied', 'count', 'Funnel'),
    ('responded', 'Responded', 'count', 'Funnel'),
    ('followed_up', 'Followed up', 'count', 'Funnel'),
    ('interviews', 'Interview / progressing', 'count', 'Funnel'),
    ('rejected', 'Rejected', 'count', 'Funnel'),
    ('accepted', 'Accepted / engaged', 'count', 'Funnel'),
    ('tokens_consumed', 'Tokens consumed', 'tokens', 'Resources'),
    ('data_downloaded', 'Data downloaded', 'bytes', 'Resources'),
]
METRIC_META = {key: {'key': key, 'label': label, 'unit': unit, 'group': group} for key, label, unit, group in METRICS}


@dataclass
class Bucket:
    start: datetime
    end: datetime
    label: str


def period_bounds(name: str, now=None):
    now = now or timezone.now()
    if name == 'today':
        # Backward-compatible query key; ScoutBox 0.8.87 makes this a true rolling window.
        return now - timedelta(hours=24), now
    if name == '3d':
        return now - timedelta(hours=72), now
    if name == 'week':
        start = now.replace(hour=0, minute=0, second=0, microsecond=0) - timedelta(days=now.weekday())
        return start, now
    if name == 'month':
        return now.replace(day=1, hour=0, minute=0, second=0, microsecond=0), now
    if name == 'quarter':
        month = ((now.month - 1) // 3) * 3 + 1
        return now.replace(month=month, day=1, hour=0, minute=0, second=0, microsecond=0), now
    if name == 'year':
        return now.replace(month=1, day=1, hour=0, minute=0, second=0, microsecond=0), now
    return None, now


def _floor_day(value):
    return value.replace(hour=0, minute=0, second=0, microsecond=0)


def _next_month(value):
    year = value.year + (1 if value.month == 12 else 0)
    month = 1 if value.month == 12 else value.month + 1
    return value.replace(year=year, month=month, day=1, hour=0, minute=0, second=0, microsecond=0)


def make_buckets(period, start, end):
    now = timezone.now()
    end = end or now
    if not start:
        # All-time without a reliable beginning is capped to the preceding 12 months for charts.
        start = end - timedelta(days=365)
    span = max(timedelta(hours=1), end - start)
    buckets = []
    if period == 'today' or span <= timedelta(days=2):
        cursor = start.replace(minute=0, second=0, microsecond=0)
        while cursor < end:
            nxt = cursor + timedelta(hours=1)
            buckets.append(Bucket(cursor, min(nxt, end), cursor.strftime('%H:%M')))
            cursor = nxt
    elif period == '3d':
        # A true rolling 72-hour window, summarized into twelve readable six-hour buckets.
        cursor = start
        while cursor < end:
            nxt = cursor + timedelta(hours=6)
            buckets.append(Bucket(cursor, min(nxt, end), cursor.strftime('%d %b %H:%M')))
            cursor = nxt
    elif period in ('week', 'month') or span <= timedelta(days=45):
        cursor = _floor_day(start)
        while cursor < end:
            nxt = cursor + timedelta(days=1)
            label = cursor.strftime('%a %d') if span <= timedelta(days=10) else cursor.strftime('%d %b')
            buckets.append(Bucket(cursor, min(nxt, end), label))
            cursor = nxt
    elif period == 'quarter' or span <= timedelta(days=150):
        cursor = _floor_day(start)
        while cursor < end:
            nxt = cursor + timedelta(days=7)
            buckets.append(Bucket(cursor, min(nxt, end), cursor.strftime('%d %b')))
            cursor = nxt
    else:
        cursor = start.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        while cursor < end:
            nxt = _next_month(cursor)
            buckets.append(Bucket(cursor, min(nxt, end), cursor.strftime('%b %Y')))
            cursor = nxt
    # Keep charts bounded for pathological custom ranges.
    if len(buckets) > 60:
        step = max(1, len(buckets) // 48)
        compact = []
        for i in range(0, len(buckets), step):
            group = buckets[i:i + step]
            compact.append(Bucket(group[0].start, group[-1].end, group[0].label))
        buckets = compact
    return buckets


def _matches_follow_up(subject, body):
    text = f'{subject or ""}\n{body or ""}'.lower()
    return bool(re.search(r'follow[ -]?up|checking in|just checking|any update|touching base|circling back', text))


def _opportunity_q(q):
    return Q(title__icontains=q) | Q(company__icontains=q) | Q(source__name__icontains=q) | Q(campaigns__name__icontains=q)


def _application_q(q):
    return (Q(opportunity__title__icontains=q) | Q(opportunity__company__icontains=q)
            | Q(opportunity__source__name__icontains=q) | Q(opportunity__campaigns__name__icontains=q))


def _bucket_index(buckets, value):
    if not value:
        return None
    for idx, bucket in enumerate(buckets):
        if bucket.start <= value < bucket.end:
            return idx
    return None


def _empty_rows(buckets):
    return [dict(bucket=b.label, **{k: 0 for k, *_ in METRICS}) for b in buckets]


def performance_series(start, end, period='month', q=''):
    """Aggregate the selected reporting interval into chart-ready buckets.

    Opportunity/application metrics respect the list search query. Infrastructure/resource
    metrics remain global because search-provider/usage records are not role-addressable.
    Performance Lab usage is excluded from production resource statistics.
    """
    end = end or timezone.now()
    buckets = make_buckets(period, start, end)
    if not buckets:
        return [], {k: 0 for k, *_ in METRICS}
    range_start, range_end = buckets[0].start, buckets[-1].end
    rows = _empty_rows(buckets)

    oqs = Opportunity.objects.filter(created_at__gte=range_start, created_at__lt=range_end, suppressed=False, user_deleted=False)
    if q:
        oqs = oqs.filter(_opportunity_q(q)).distinct()
    for created_at, fit_score, status in oqs.values_list('created_at', 'fit_score', 'status').iterator():
        idx = _bucket_index(buckets, created_at)
        if idx is None:
            continue
        rows[idx]['roles_found'] += 1
        if status == 'apply' or int(fit_score or 0) >= 80:
            rows[idx]['high_priority'] += 1

    aqs = Application.objects.filter(deleted_at__isnull=True, opportunity__isnull=False, opportunity__user_deleted=False)
    if q:
        aqs = aqs.filter(_application_q(q)).distinct()
    for app_id, created_at in aqs.filter(created_at__gte=range_start, created_at__lt=range_end, opportunity__suppressed=False).values_list('id', 'created_at').iterator():
        idx = _bucket_index(buckets, created_at)
        if idx is not None:
            rows[idx]['prepared'] += 1
    for app_id, applied_at in aqs.filter(applied_at__gte=range_start, applied_at__lt=range_end).values_list('id', 'applied_at').iterator():
        idx = _bucket_index(buckets, applied_at)
        if idx is not None:
            rows[idx]['newly_applied'] += 1

    # Incoming mail is the strongest timestamped evidence of an actual response.
    mqs = MailEvent.objects.filter(application__isnull=False, application__opportunity__user_deleted=False, occurred_at__gte=range_start, occurred_at__lt=range_end)
    if q:
        mqs = mqs.filter(
            Q(application__opportunity__title__icontains=q)
            | Q(application__opportunity__company__icontains=q)
            | Q(application__opportunity__source__name__icontains=q)
            | Q(application__opportunity__campaigns__name__icontains=q)
        ).distinct()
    responded_seen = defaultdict(set)
    follow_seen = defaultdict(set)
    rejected_seen = defaultdict(set)
    interview_seen = defaultdict(set)
    accepted_seen = defaultdict(set)
    for app_id, occurred_at, kind, classification, subject, body in mqs.values_list(
        'application_id', 'occurred_at', 'kind', 'classification', 'subject', 'body_excerpt'
    ).iterator():
        idx = _bucket_index(buckets, occurred_at)
        if idx is None:
            continue
        if kind == 'inbox':
            responded_seen[idx].add(app_id)
            if classification == 'rejected':
                rejected_seen[idx].add(app_id)
            elif classification == 'interview':
                interview_seen[idx].add(app_id)
            elif classification == 'accepted':
                accepted_seen[idx].add(app_id)
        elif kind == 'sent' and _matches_follow_up(subject, body):
            follow_seen[idx].add(app_id)

    # Status changes without a timestamped classification email (manual/imported updates).
    for status, target in [('rejected', rejected_seen), ('interview', interview_seen), ('accepted', accepted_seen)]:
        sqs = aqs.filter(status=status, opportunity__suppressed=False, updated_at__gte=range_start, updated_at__lt=range_end)
        for app_id, updated_at in sqs.values_list('id', 'updated_at').iterator():
            idx = _bucket_index(buckets, updated_at)
            if idx is not None:
                target[idx].add(app_id)

    for idx in range(len(rows)):
        rows[idx]['responded'] = len(responded_seen[idx])
        rows[idx]['followed_up'] = len(follow_seen[idx])
        rows[idx]['interviews'] = len(interview_seen[idx])
        rows[idx]['rejected'] = len(rejected_seen[idx])
        rows[idx]['accepted'] = len(accepted_seen[idx])

    # Provider stats are daily aggregates and intentionally stay global even when a role search is active.
    last_day = (range_end - timedelta(microseconds=1)).date()
    pqs = SearchProviderStat.objects.filter(day__gte=range_start.date(), day__lte=last_day)
    for day, duplicates, applied_matches in pqs.values_list('day', 'duplicates', 'applied_matches').iterator():
        dt = timezone.make_aware(datetime.combine(day, datetime.min.time()), timezone.get_current_timezone())
        idx = _bucket_index(buckets, dt)
        if idx is not None:
            rows[idx]['duplicates'] += int(duplicates or 0)
            rows[idx]['already_applied'] += int(applied_matches or 0)

    uqs = UsageMetric.objects.filter(at__gte=range_start, at__lt=range_end).exclude(category='lab')
    for at, category, stage, pages, tin, tout, b in uqs.values_list(
        'at', 'category', 'stage', 'pages', 'tokens_in', 'tokens_out', 'bytes_downloaded'
    ).iterator():
        idx = _bucket_index(buckets, at)
        if idx is None:
            continue
        if category == 'scrape' and (stage == 'page_fetch' or int(pages or 0) > 0):
            rows[idx]['pages_scraped'] += int(pages or 0)
        rows[idx]['tokens_consumed'] += int(tin or 0) + int(tout or 0)
        rows[idx]['data_downloaded'] += int(b or 0)

    totals = {k: sum(int(row[k] or 0) for row in rows) for k, *_ in METRICS}
    return rows, totals


def performance_totals(start, end, q=''):
    end = end or timezone.now()
    totals = {k: 0 for k, *_ in METRICS}
    oqs = Opportunity.objects.filter(created_at__gte=start, created_at__lt=end, suppressed=False, user_deleted=False)
    if q:
        oqs = oqs.filter(_opportunity_q(q)).distinct()
    totals['roles_found'] = oqs.count()
    totals['high_priority'] = oqs.filter(Q(status='apply') | Q(fit_score__gte=80)).count()

    aqs = Application.objects.filter(deleted_at__isnull=True, opportunity__isnull=False, opportunity__user_deleted=False)
    if q:
        aqs = aqs.filter(_application_q(q)).distinct()
    totals['prepared'] = aqs.filter(created_at__gte=start, created_at__lt=end, opportunity__suppressed=False).count()
    totals['newly_applied'] = aqs.filter(applied_at__gte=start, applied_at__lt=end).count()

    mqs = MailEvent.objects.filter(application__isnull=False, application__opportunity__user_deleted=False, occurred_at__gte=start, occurred_at__lt=end)
    if q:
        mqs = mqs.filter(
            Q(application__opportunity__title__icontains=q)
            | Q(application__opportunity__company__icontains=q)
            | Q(application__opportunity__source__name__icontains=q)
            | Q(application__opportunity__campaigns__name__icontains=q)
        ).distinct()
    totals['responded'] = mqs.filter(kind='inbox').values('application_id').distinct().count()
    follow_ids = set()
    for app_id, subject, body in mqs.filter(kind='sent').values_list('application_id', 'subject', 'body_excerpt').iterator():
        if _matches_follow_up(subject, body):
            follow_ids.add(app_id)
    totals['followed_up'] = len(follow_ids)

    classification_sets = {}
    for classification, status, key in [('interview', 'interview', 'interviews'), ('rejected', 'rejected', 'rejected'), ('accepted', 'accepted', 'accepted')]:
        ids = set(mqs.filter(kind='inbox', classification=classification).values_list('application_id', flat=True).distinct())
        ids.update(aqs.filter(status=status, opportunity__suppressed=False, updated_at__gte=start, updated_at__lt=end).values_list('id', flat=True))
        classification_sets[key] = ids
        totals[key] = len(ids)

    last_day = (end - timedelta(microseconds=1)).date()
    p = SearchProviderStat.objects.filter(day__gte=start.date(), day__lte=last_day)
    p_sum = p.aggregate(duplicates=Sum('duplicates'), applied=Sum('applied_matches'))
    totals['duplicates'] = int(p_sum.get('duplicates') or 0)
    totals['already_applied'] = int(p_sum.get('applied') or 0)

    uqs = UsageMetric.objects.filter(at__gte=start, at__lt=end).exclude(category='lab')
    u = uqs.aggregate(tokens_in=Sum('tokens_in'), tokens_out=Sum('tokens_out'), bytes=Sum('bytes_downloaded'))
    totals['tokens_consumed'] = int(u.get('tokens_in') or 0) + int(u.get('tokens_out') or 0)
    totals['data_downloaded'] = int(u.get('bytes') or 0)
    pages = uqs.filter(category='scrape').aggregate(pages=Sum('pages'))
    totals['pages_scraped'] = int(pages.get('pages') or 0)
    return totals


def comparison_periods(q=''):
    now = timezone.now()
    out = {}
    for name, label in [('week', 'This week'), ('month', 'This month'), ('quarter', 'This quarter'), ('year', 'This year')]:
        start, end = period_bounds(name, now)
        totals = performance_totals(start, end, q=q)
        out[name] = {'label': label, **totals}
    return out


def human_bytes(value):
    n = float(value or 0)
    for unit in ('B', 'KB', 'MB', 'GB', 'TB'):
        if n < 1024 or unit == 'TB':
            return f'{int(n)} B' if unit == 'B' else f'{n:.1f} {unit}'
        n /= 1024


def export_csv(series, totals, comparison, period_label):
    import csv
    import io
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(['ScoutBox Statistics', period_label])
    w.writerow([])
    w.writerow(['Metric', 'Value'])
    for key, label, unit, _ in METRICS:
        w.writerow([label, totals.get(key, 0)])
    w.writerow([])
    w.writerow(['Metric', 'This week', 'This month', 'This quarter', 'This year'])
    for key, label, _, _ in METRICS:
        w.writerow([label, comparison['week'][key], comparison['month'][key], comparison['quarter'][key], comparison['year'][key]])
    w.writerow([])
    headers = ['Bucket'] + [label for _, label, _, _ in METRICS]
    w.writerow(headers)
    for row in series:
        w.writerow([row['bucket']] + [row[k] for k, *_ in METRICS])
    return buf.getvalue().encode('utf-8-sig')


def export_xlsx(series, totals, comparison, period_label):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Selected period'
    ws.append(['ScoutBox Statistics', period_label])
    ws.append([])
    ws.append(['Metric', 'Value'])
    for key, label, _, _ in METRICS:
        ws.append([label, totals.get(key, 0)])
    ws['A1'].font = Font(bold=True, size=14)
    ws['A3'].font = ws['B3'].font = Font(bold=True)
    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['B'].width = 18

    trend = wb.create_sheet('Trend data')
    trend.append(['Bucket'] + [label for _, label, _, _ in METRICS])
    for row in series:
        trend.append([row['bucket']] + [row[k] for k, *_ in METRICS])
    for cell in trend[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(bold=True, color='FFFFFF')
    trend.freeze_panes = 'A2'
    trend.column_dimensions['A'].width = 16
    for col in range(2, len(METRICS) + 2):
        trend.column_dimensions[get_column_letter(col)].width = 18

    # Native Excel charts give the downloaded workbook a useful visual summary too.
    if len(series) > 1:
        disc = LineChart()
        disc.title = 'Discovery activity'
        disc.y_axis.title = 'Count'
        disc.x_axis.title = 'Period'
        disc.height = 8
        disc.width = 15
        data = Reference(trend, min_col=2, max_col=5, min_row=1, max_row=len(series) + 1)
        cats = Reference(trend, min_col=1, min_row=2, max_row=len(series) + 1)
        disc.add_data(data, titles_from_data=True)
        disc.set_categories(cats)
        ws.add_chart(disc, 'D3')

        funnel = LineChart()
        funnel.title = 'Funnel movement'
        funnel.y_axis.title = 'Count'
        funnel.x_axis.title = 'Period'
        funnel.height = 8
        funnel.width = 15
        data2 = Reference(trend, min_col=6, max_col=13, min_row=1, max_row=len(series) + 1)
        funnel.add_data(data2, titles_from_data=True)
        funnel.set_categories(cats)
        ws.add_chart(funnel, 'D20')

    comp = wb.create_sheet('Week Month Quarter Year')
    comp.append(['Metric', 'This week', 'This month', 'This quarter', 'This year'])
    for key, label, _, _ in METRICS:
        comp.append([label, comparison['week'][key], comparison['month'][key], comparison['quarter'][key], comparison['year'][key]])
    for cell in comp[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='17365D')
    comp.freeze_panes = 'A2'
    comp.column_dimensions['A'].width = 26
    for c in ('B', 'C', 'D', 'E'):
        comp.column_dimensions[c].width = 18

    defs = wb.create_sheet('Definitions')
    defs.append(['Metric', 'Definition'])
    definitions = metric_definitions()
    for key, label, _, _ in METRICS:
        defs.append([label, definitions[key]])
    defs.column_dimensions['A'].width = 24
    defs.column_dimensions['B'].width = 100

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _pdf_chart(series, keys, title, width=500, height=170, labels=None):
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.charts.legends import Legend
    from reportlab.lib import colors

    drawing = Drawing(width, height)
    drawing.add(String(0, height - 14, title, fontSize=10, fillColor=colors.HexColor('#17365D')))
    chart = HorizontalLineChart()
    chart.x = 35
    chart.y = 35
    chart.height = height - 65
    chart.width = width - 55
    chart.data = [[row[k] for row in series] for k in keys]
    chart.categoryAxis.categoryNames = [row['bucket'] for row in series]
    chart.categoryAxis.labels.fontSize = 6
    chart.categoryAxis.labels.angle = 25 if len(series) > 10 else 0
    chart.valueAxis.valueMin = 0
    chart.valueAxis.valueMax = max(1, max((max(values) if values else 0) for values in chart.data))
    chart.valueAxis.labels.fontSize = 7
    palette = [colors.HexColor(x) for x in ('#167FB1', '#21A366', '#B77B26', '#7656A5', '#C94F5A')]
    for idx in range(len(keys)):
        chart.lines[idx].strokeColor = palette[idx % len(palette)]
        chart.lines[idx].strokeWidth = 1.5
        chart.lines[idx].symbol = None
    drawing.add(chart)
    legend = Legend()
    legend.x = 38
    legend.y = 12
    legend.fontSize = 6
    legend.dxTextSpace = 3
    legend.columnMaximum = 3
    labels = labels or {k: METRIC_META.get(k, {}).get('label', k) for k in keys}
    legend.colorNamePairs = [(palette[i % len(palette)], labels.get(k, k)) for i, k in enumerate(keys)]
    drawing.add(legend)
    return drawing


def export_pdf(series, totals, comparison, period_label):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=10 * mm, bottomMargin=10 * mm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Small', parent=styles['BodyText'], fontSize=7.5, leading=9, alignment=TA_LEFT))
    story = [
        Paragraph('ScoutBox — Statistics', styles['Title']),
        Paragraph(period_label, styles['Heading3']),
        Paragraph('Production resource totals exclude Performance Lab benchmark traffic. Opportunity/application metrics honor the active statistics search filter; provider/resource totals remain system-wide.', styles['Small']),
        Spacer(1, 5 * mm),
    ]
    summary = [['Metric', 'Value']]
    for key, label, unit, _ in METRICS:
        val = totals.get(key, 0)
        if unit == 'bytes': val = human_bytes(val)
        elif unit == 'tokens': val = f'{int(val):,}'
        summary.append([label, val])
    table = Table(summary, colWidths=[55 * mm, 36 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#17365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#AAB7C4')),
        ('FONTSIZE', (0, 0), (-1, -1), 7.5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F6F9')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    story.append(Spacer(1, 5 * mm))
    if series:
        story.append(_pdf_chart(series, ['pages_scraped', 'roles_found', 'duplicates', 'already_applied'], 'Discovery activity'))
        story.append(Spacer(1, 3 * mm))
        story.append(_pdf_chart(series, ['high_priority', 'newly_applied', 'responded', 'followed_up', 'rejected'], 'Funnel movement'))
        story.append(PageBreak())
        story.append(_pdf_chart(series, ['tokens_consumed'], 'Production AI tokens consumed', height=150))
        # Bytes get their own chart to avoid mixing incompatible scales.
        data_series = [dict(row, data_downloaded_mb=round(row['data_downloaded'] / (1024 * 1024), 3)) for row in series]
        story.append(_pdf_chart(data_series, ['data_downloaded_mb'], 'Data downloaded (MB)', height=150, labels={'data_downloaded_mb': 'Downloaded (MB)'}))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph('Week / month / quarter / year comparison', styles['Heading3']))
    comp_rows = [['Metric', 'This week', 'This month', 'This quarter', 'This year']]
    for key, label, unit, _ in METRICS:
        vals = [comparison[p][key] for p in ('week', 'month', 'quarter', 'year')]
        if unit == 'bytes': vals = [human_bytes(v) for v in vals]
        elif unit == 'tokens': vals = [f'{int(v):,}' for v in vals]
        comp_rows.append([label] + vals)
    comp_table = Table(comp_rows, repeatRows=1, colWidths=[55 * mm, 35 * mm, 35 * mm, 35 * mm, 35 * mm])
    comp_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#17365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#AAB7C4')),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F6F9')]),
    ]))
    story.append(comp_table)
    doc.build(story)
    return out.getvalue()


def metric_definitions():
    return {
        'pages_scraped': 'Production page-fetch events recorded by ScoutBox. Performance Lab scrape tests are excluded.',
        'roles_found': 'New Opportunity records first created during the period after discovery consolidation.',
        'duplicates': 'Search results suppressed as duplicates by provider/day statistics.',
        'already_applied': 'Search results matched to previously applied roles and suppressed/downranked before normal processing.',
        'high_priority': 'Newly found opportunities marked Apply Now or carrying a fit score of at least 80.',
        'prepared': 'Application preparation records created for non-suppressed opportunities during the period.',
        'newly_applied': 'Applications whose applied_at timestamp falls inside the period.',
        'responded': 'Distinct applications with at least one inbound IMAP message observed in the period.',
        'followed_up': 'Distinct applications with an observed Sent message that looks like a follow-up/check-in.',
        'interviews': 'Distinct applications with interview/progression mail or a corresponding status update during the period.',
        'rejected': 'Distinct applications with rejection mail or a rejection status update during the period.',
        'accepted': 'Distinct applications with acceptance/engagement mail or a corresponding status update during the period.',
        'tokens_consumed': 'Input plus output tokens recorded for production AI stages; Performance Lab runs are excluded.',
        'data_downloaded': 'Response bytes recorded in production usage telemetry, including search/page-fetch traffic; Performance Lab is excluded.',
    }


def earliest_activity(q=''):
    """Best-effort beginning for an all-time chart/export."""
    candidates = []
    oqs = Opportunity.objects.filter(user_deleted=False)
    if q:
        oqs = oqs.filter(_opportunity_q(q)).distinct()
    first = oqs.order_by('created_at').values_list('created_at', flat=True).first()
    if first:
        candidates.append(first)
    aqs = Application.objects.filter(opportunity__user_deleted=False,deleted_at__isnull=True).exclude(applied_at=None)
    if q:
        aqs = aqs.filter(_application_q(q)).distinct()
    first = aqs.order_by('applied_at').values_list('applied_at', flat=True).first()
    if first:
        candidates.append(first)
    first = UsageMetric.objects.exclude(category='lab').order_by('at').values_list('at', flat=True).first()
    if first:
        candidates.append(first)
    day = SearchProviderStat.objects.order_by('day').values_list('day', flat=True).first()
    if day:
        candidates.append(timezone.make_aware(datetime.combine(day, datetime.min.time()), timezone.get_current_timezone()))
    return min(candidates) if candidates else timezone.now() - timedelta(days=365)
