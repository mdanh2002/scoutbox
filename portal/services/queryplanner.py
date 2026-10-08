"""Resume-first automatic search query planning.

The planner builds profession-neutral role/skill queries. Discovery-market geography is
added later by the acquisition layer; Candidate Profile operating locations remain a
soft ranking/suitability signal and do not suppress initial market searches.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import time
from collections import Counter
from typing import Iterable

from portal.models import DocumentAsset, Profile
from .source_domains import expand_source_domains


# Candidate Profile search vocabulary is intentionally generous. The saved profile is
# an editable summary of Resume evidence, not a narrow whitelist. These limits protect
# the UI/database from accidental unbounded input while leaving enough headroom for a
# broad multi-disciplinary CV.
PROFILE_CONCEPT_LIMIT = 160
PROFILE_ROLE_LIMIT = 80


# The graph is deliberately broader than literal CV vocabulary.  "aliases" are
# reasonable search adjacencies, not assertions that two concepts are identical.
SKILL_GRAPH = {
    'embedded systems': {
        'patterns': ['embedded systems', 'embedded system', 'embedded linux'],
        'aliases': ['embedded software', 'microcontroller', 'MCU', 'bare metal', 'RTOS'],
        'category': 'embedded', 'weight': 10,
    },
    'microcontroller': {
        'patterns': ['microcontroller', 'pic32', 'pic24', 'dspic', 'stm32', 'esp32', 'avr', 'raspberry pi'],
        'aliases': ['MCU', 'embedded firmware', 'bare metal', 'PIC', 'STM32', 'ESP32'],
        'category': 'embedded', 'weight': 9,
    },
    'firmware': {
        'patterns': ['firmware', 'u-boot', 'bootloader', 'boot image', 'rom image', 'option rom'],
        'aliases': ['embedded firmware', 'bootloader', 'U-Boot', 'ROM analysis', 'firmware engineer'],
        'category': 'embedded', 'weight': 10,
    },
    'device drivers': {
        'patterns': ['device driver', 'kernel driver', 'linux driver', 'windows driver'],
        'aliases': ['kernel driver', 'Linux device driver', 'driver development', 'driver analysis'],
        'category': 'systems', 'weight': 9,
    },
    'reverse engineering': {
        'patterns': ['reverse engineering', 'reverse engineered', 'disassembled', 'disassembly'],
        'aliases': ['binary analysis', 'firmware analysis', 'driver analysis', 'protocol analysis', 'static analysis', 'dynamic analysis'],
        'category': 'reverse', 'weight': 12,
    },
    'binary analysis': {
        'patterns': ['binary analysis', 'ghidra', 'ida pro', 'radare2', 'windbg', 'frida'],
        'aliases': ['binary reverse engineering', 'static analysis', 'dynamic analysis', 'disassembly'],
        'category': 'reverse', 'weight': 10,
    },
    'protocol reverse engineering': {
        'patterns': ['protocol reverse engineering', 'proprietary protocol', 'protocol analysis', 'wireshark'],
        'aliases': ['protocol analysis', 'protocol reverse engineering', 'undocumented protocol', 'packet analysis'],
        'category': 'reverse', 'weight': 10,
    },
    'QEMU': {
        'patterns': ['qemu', 'qemu/kvm', 'qmp', 'qcow2', 'virtio'],
        'aliases': ['emulation', 'virtual device', 'device model', 'virtio', 'KVM', 'QMP'],
        'category': 'virtualization', 'weight': 13,
    },
    'virtualization': {
        'patterns': ['virtualization', 'hypervisor', 'firecracker', 'gvisor', 'kata containers', 'libvirt', 'kvm'],
        'aliases': ['hypervisor', 'virtual machine monitor', 'KVM', 'Firecracker', 'Kata Containers'],
        'category': 'virtualization', 'weight': 9,
    },
    'emulation': {
        'patterns': ['emulation', 'emulator', 'dosbox-x', '86box', 'pcem', 'mame', 'retroarch', 'bochs', 'pcjs'],
        'aliases': ['emulator development', 'hardware emulation', 'CPU emulation', 'device emulation'],
        'category': 'retro', 'weight': 12,
    },
    'retro computing': {
        'patterns': ['retro computing', 'retro-computing', 'classic pc', 'classic 68k', 'legacy hardware', 'vintage hardware'],
        'aliases': ['legacy systems', 'vintage computing', 'classic hardware', 'retro software'],
        'category': 'retro', 'weight': 11,
    },
    'x86': {
        'patterns': ['x86', '8086', 'pc xt', 'real mode'],
        'aliases': ['x86 systems', '8086', 'real mode', 'legacy PC'],
        'category': 'retro', 'weight': 10,
    },
    'BIOS': {
        'patterns': ['bios', 'int 13h', 'option rom'],
        'aliases': ['BIOS engineering', 'ROM analysis', 'option ROM', 'legacy BIOS'],
        'category': 'retro', 'weight': 11,
    },
    'DOS': {
        'patterns': ['dos', 'ms-dos', 'dosbox'],
        'aliases': ['MS-DOS', 'DOS compatibility', 'legacy DOS', 'DOS emulator'],
        'category': 'retro', 'weight': 9,
    },
    'C/C++': {
        'patterns': ['c/c++', 'c++', 'systems programming'],
        'aliases': ['C++ systems', 'systems programming', 'low-level C++'],
        'category': 'systems', 'weight': 9,
    },
    'Linux systems': {
        'patterns': ['linux internals', 'embedded linux', 'linux systems', 'debian', 'buildroot', 'yocto'],
        'aliases': ['Linux systems', 'Linux internals', 'embedded Linux', 'Buildroot', 'Yocto'],
        'category': 'systems', 'weight': 9,
    },
    'PIC / dsPIC': {
        'patterns': ['pic32', 'pic32mz', 'pic24', 'dspic', 'dspic33'],
        'aliases': ['Microchip PIC', 'PIC firmware', 'PIC microcontroller', 'dsPIC firmware'],
        'category': 'embedded-niche', 'weight': 10,
    },
    'STM32 / ESP32': {
        'patterns': ['stm32', 'esp32'],
        'aliases': ['ARM microcontroller', 'STM32 firmware', 'ESP32 firmware'],
        'category': 'embedded-niche', 'weight': 8,
    },
    'legacy peripheral emulation': {
        'patterns': ['zip100', 'zip drive', 'ch375', 'ch376', 'ne2000', 'isa-to-usb', 'parallel port', 'lpt protocol'],
        'aliases': ['legacy peripheral', 'ISA hardware', 'parallel port device', 'SCSI device emulation', 'legacy storage emulation'],
        'category': 'retro-niche', 'weight': 12,
    },
    'DOSBox-X': {
        'patterns': ['dosbox-x', 'dosbox'],
        'aliases': ['DOSBox', 'DOS emulator', 'PC compatibility emulator'],
        'category': 'retro-niche', 'weight': 10,
    },
    'Buildroot / Yocto': {
        'patterns': ['buildroot', 'yocto'],
        'aliases': ['embedded Linux build system', 'custom Linux image', 'embedded Linux platform'],
        'category': 'embedded-niche', 'weight': 8,
    },
    'BACnet / Modbus': {
        'patterns': ['bacnet', 'modbus'],
        'aliases': ['building automation protocol', 'industrial protocol', 'fieldbus integration'],
        'category': 'protocol-niche', 'weight': 6,
    },
    'building automation': {
        'patterns': ['building automation', 'building management system', 'building management systems', 'bms', 'bas'],
        'aliases': ['building automation systems', 'BMS', 'BAS', 'building controls', 'automation controls'],
        'category': 'protocol-niche', 'weight': 9,
    },
    'software integration': {
        'patterns': ['saas integration', 'rest api', 'rest apis', 'soap', 'graphql', 'mulesoft'],
        'aliases': ['systems integration', 'API integration', 'enterprise integration', 'integration engineering'],
        'category': 'integration', 'weight': 7,
    },
    'cloud platforms': {
        'patterns': ['gcp', 'azure', 'aws', 'google cloud', 'amazon web services'],
        'aliases': ['cloud engineering', 'cloud platform', 'cloud infrastructure'],
        'category': 'cloud', 'weight': 6,
    },
    'data engineering': {
        'patterns': ['azure data factory', 'aws glue', 'dbt', 'apache spark', 'etl', 'ssis'],
        'aliases': ['data pipeline', 'ETL', 'data integration', 'analytics engineering'],
        'category': 'data', 'weight': 6,
    },
    'hardware interfaces': {
        'patterns': ['spi', 'i²c', 'i2c', 'uart', 'jtag', 'can', 'gpio', 'rs-485', 'rs485'],
        'aliases': ['hardware interface', 'board bring-up', 'JTAG', 'serial protocol'],
        'category': 'embedded', 'weight': 7,
    },
    'network protocols': {
        'patterns': ['tcp/ip', 'udp', 'ethernet', 'rtsp', 'sip', 'bacnet', 'modbus', 'onvif'],
        'aliases': ['network protocol', 'industrial protocol', 'protocol integration', 'device networking'],
        'category': 'systems', 'weight': 7,
    },
    'technical writing': {
        'patterns': ['technical writer', 'technical writing', 'product documentation', 'user guides', 'help-center', 'documentation workflows'],
        'aliases': ['documentation engineer', 'developer documentation', 'technical content', 'technical author'],
        'category': 'writing', 'weight': 10,
    },
    'developer education': {
        'patterns': ['developer education', 'educator', 'teaching', 'training', 'instructor'],
        'aliases': ['developer educator', 'technical trainer', 'technical instructor', 'developer advocate'],
        'category': 'writing', 'weight': 7,
    },
    'security/TLS': {
        'patterns': ['mbedtls', 'wolfssl', 'openssl', 'tls 1.2', 'tls 1.3', 'x.509', 'certificate chains'],
        'aliases': ['embedded security', 'TLS', 'X.509', 'secure communications'],
        'category': 'security', 'weight': 7,
    },
    'Python tooling': {
        'patterns': ['python', 'pandas', 'matplotlib'],
        'aliases': ['Python tooling', 'automation', 'engineering tools'],
        'category': 'tooling', 'weight': 5,
    },
}

ROLE_FAMILIES = {
    'embedded software engineer': ['embedded systems', 'microcontroller', 'firmware', 'hardware interfaces'],
    'firmware engineer': ['firmware', 'microcontroller', 'device drivers'],
    'reverse engineer': ['reverse engineering', 'binary analysis', 'protocol reverse engineering'],
    'systems software engineer': ['C/C++', 'Linux systems', 'device drivers', 'QEMU'],
    'emulation engineer': ['emulation', 'QEMU', 'x86', 'BIOS', 'DOS'],
    'virtualization engineer': ['QEMU', 'virtualization', 'Linux systems'],
    'legacy systems specialist': ['retro computing', 'x86', 'BIOS', 'DOS', 'emulation'],
    'building automation software engineer': ['building automation', 'BACnet / Modbus', 'network protocols', 'software integration'],
    'building automation engineer': ['building automation', 'BACnet / Modbus', 'network protocols'],
    'industrial protocol integration engineer': ['BACnet / Modbus', 'network protocols', 'protocol reverse engineering', 'software integration'],
    'software integration engineer': ['software integration', 'network protocols', 'cloud platforms'],
    'cloud integration engineer': ['cloud platforms', 'software integration', 'Docker', 'Kubernetes', 'Terraform'],
    'data engineer': ['data engineering', 'PostgreSQL', 'SQL Server', 'Apache Spark', 'dbt'],
    'devops platform engineer': ['cloud platforms', 'Linux systems', 'Docker', 'Kubernetes', 'Terraform', 'CI/CD'],
    'full stack software engineer': ['Java', 'Python', 'C/C++', 'Node.js', 'React', 'Angular', 'ASP.NET'],
    'technical writer': ['technical writing', 'developer education'],
    'technical trainer': ['developer education', 'technical writing', 'embedded systems'],
}


PREFERENCE_TECH_HINTS = {
    'embedded','firmware','microcontroller','mcu','qemu','virtualization','emulation','reverse','engineering','binary','protocol','bios','dos','x86','legacy','retro','driver','kernel','linux','rtos','fpga','pic','stm32','esp32','writing','documentation','technical','teaching','training','security','tls','hardware','systems','c++','c/c++'
}


def _explicit_preference_phrases(text: str) -> list[str]:
    """Extract short explicit preference phrases not already covered by the skill graph.

    Free-text preference boxes remain useful without turning arbitrary prose into a giant
    keyword bag.  Only concise, technical-looking comma/semicolon/newline segments are
    admitted as search steering terms.
    """
    out=[]
    for raw in re.split(r'[,;\n]+', text or ''):
        phrase=_normalise(raw).strip(' .:-')
        low=phrase.lower()
        if not phrase or any(x in low for x in ['avoid ', 'exclude ', 'reject ', 'not interested', 'do not ']):
            continue
        phrase=re.sub(r'(?i)^(prefer|priority|interested in|looking for)\s+','',phrase).strip()
        words=re.findall(r'[A-Za-z0-9+#/.-]+',phrase)
        if not 1 <= len(words) <= 6:
            continue
        if not any(w.lower() in PREFERENCE_TECH_HINTS or re.search(r'[A-Z0-9+#/.-]',w) for w in words):
            continue
        # Drop vague trailing nouns that add noise without changing the technical concept.
        while words and words[-1].lower() in {'role','roles','job','jobs','opportunity','opportunities','work'}:
            words.pop()
        if words:
            candidate=' '.join(words)
            if candidate.lower() not in {x.lower() for x in out}: out.append(candidate)
    return out[:20]

LABEL_FOCUS_HINTS = {
    'technical writing': ['technical writing','developer education'],
    'embedded': ['embedded systems','microcontroller','firmware','PIC / dsPIC','STM32 / ESP32'],
    'retro': ['retro computing','emulation','legacy peripheral emulation','x86','BIOS','DOS'],
    'reverse engineer': ['reverse engineering','binary analysis','protocol reverse engineering'],
    'virtualization': ['QEMU','virtualization','emulation'],
    'software engineer': ['C/C++','Linux systems','Python tooling'],
}

NEGATIVE_HINTS = {
    'onsite': ['onsite', 'on-site', 'office presence'],
    'leetcode': ['leetcode', 'live coding'],
    'psychometric': ['psychometric'],
}


def _normalise(text: str) -> str:
    return re.sub(r'\s+', ' ', (text or '').replace('\u00b2', '2')).strip()


def _has(text: str, phrase: str) -> bool:
    text = text.lower()
    phrase = phrase.lower().strip()
    if not phrase:
        return False
    # Word boundaries are unhelpful for C++, I2C and similar punctuation-heavy terms.
    if any(ch in phrase for ch in '+/#².-'):
        return phrase in text
    return re.search(r'(?<![a-z0-9])' + re.escape(phrase) + r'(?![a-z0-9])', text) is not None


def _read_asset_text(asset: DocumentAsset) -> str:
    """Extract searchable text from one active Resume without altering the source file."""
    name = (asset.original_name or asset.file.name or '').lower()
    try:
        with asset.file.open('rb') as fh:
            payload = fh.read()
        if name.endswith('.docx'):
            from docx import Document
            doc = Document(io.BytesIO(payload))
            chunks = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    chunks.append(' | '.join(c.text for c in row.cells))
            return _normalise('\n'.join(chunks))
        if name.endswith('.pdf'):
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(payload))
            return _normalise('\n'.join((p.extract_text() or '') for p in reader.pages))
    except Exception:
        return ''
    return ''


def extract_active_cv_texts() -> list[dict]:
    docs = []
    for asset in DocumentAsset.objects.filter(kind='cv', active=True).order_by('id'):
        text = _read_asset_text(asset)
        if text:
            docs.append({'id': asset.pk, 'label': asset.label, 'name': asset.original_name, 'text': text})
    return docs


def _split_config(value: str) -> list[str]:
    return [x.strip() for x in re.split(r'[,;\n|]+', value or '') if x.strip()]


def _preference_negatives(profile: Profile) -> list[str]:
    blob = ' '.join([profile.high_priority_text, profile.medium_priority_text, profile.low_priority_text]).lower()
    out = []
    for canonical, hints in NEGATIVE_HINTS.items():
        for hint in hints:
            # Only infer a negative preference when language close to the phrase signals rejection/avoidance.
            m = re.search(r'(avoid|exclude|reject|no|not interested in)[^.\n]{0,80}' + re.escape(hint), blob)
            if m:
                out.append(canonical)
                break
    return out




_GENERIC_FULL_STACK_RE = re.compile(r'(?i)\bgeneric[\s-]*full[\s-]*stack\b')
_EXTERNAL_EXCLUSION_RE = re.compile(r'(?<!\S)-(?:(?:"[^"]+")|(?:\'[^\']+\')|(?:[^\s]+))')

def strip_external_exclusions(value: str) -> str:
    """Remove engine-specific unary exclusions before a query leaves ScoutBox.

    Exclusion policy belongs in the deterministic post-retrieval filter. Several public
    search endpoints treat leading '-' literally, which can make an exclusion phrase a
    required search phrase instead of an exclusion.
    """
    text=str(value or '')
    text=_EXTERNAL_EXCLUSION_RE.sub(' ',text)
    text=_GENERIC_FULL_STACK_RE.sub(' ',text)
    return re.sub(r'\s+',' ',text).strip()



def _normalize_site_constraint(token: str) -> str:
    """Normalize site: operators to one host-only constraint.

    Search engines vary in how they interpret path-scoped operators such as
    ``site:facebook.com/posts``; some treat the path as a literal term or over-constrain
    the request. ScoutBox should use source/path knowledge after retrieval, so outgoing
    queries keep only ``site:domain.tld``.
    """
    raw=re.sub(r'[()\[\]{}]+','',str(token or '')).strip()
    if not raw.lower().startswith('site:'):
        return raw
    value=raw[5:].strip().strip('"\'')
    value=re.sub(r'^https?://','',value,flags=re.I)
    host=value.split('/',1)[0].split('?',1)[0].split('#',1)[0].split(':',1)[0].lower().removeprefix('www.')
    if not host:
        return ''
    return 'site:'+host


def _append_site_constraint_once(out: list[str], cleaned: str, seen_site: list[bool]) -> None:
    """Append at most one site: constraint to an outgoing provider query.

    A query such as ``site:facebook.com site:oracle.com "firmware engineer"`` is
    invalid/over-constrained for ScoutBox's provider rotation. This helper preserves the
    first source scope and drops later site: operators; deterministic post-retrieval
    validation still decides whether a returned page belongs to an allowed/source domain.
    ``seen_site`` is a tiny mutable box so the sanitizer loops stay compact.
    """
    if not cleaned:
        return
    if cleaned.lower().startswith('site:'):
        if seen_site[0]:
            return
        seen_site[0]=True
    out.append(cleaned)


_SEARCH_FILLER_PREPOSITIONS = {
    'in','at','near','around','within','for','with','from','to','of','by','on','into','across',
}

def _is_search_filler_word(value: str) -> bool:
    """Return True for standalone low-value search prepositions.

    Quoted phrases are never passed here, so exact user/technical phrases keep their
    wording. This is intentionally conservative and only strips whole standalone words.
    """
    token=str(value or '').strip().casefold()
    return token in _SEARCH_FILLER_PREPOSITIONS


def _repair_unbalanced_generated_quotes(value: str) -> str:
    text=str(value or '')
    if text.count(chr(34)) % 2:
        text=text.replace(chr(34),' ')
    return re.sub(r'\s+',' ',text).strip()

def sanitize_search_engine_query(value: str) -> str:
    """Keep only plain terms, quoted phrases and host-only site: constraints.

    Boolean query construction belongs in ScoutBox's post-retrieval filtering.  This
    sanitizer deliberately preserves quoted phrases and ``site:domain`` while stripping
    AND/OR/NOT tokens, pipe-style Boolean alternatives, grouping punctuation and
    accidental path-scoped site operators before any search-provider request.
    """
    text=_repair_unbalanced_generated_quotes(strip_external_exclusions(str(value or '')))
    tokens=re.findall(r'"[^"\n]*"|site:[^\s]+|[^\s]+',text.replace('|',' '),flags=re.I)
    out=[]; seen_site=[False]
    for token in tokens:
        if token.startswith('"') and token.endswith('"'):
            cleaned=token
        elif token.lower().startswith('site:'):
            cleaned=_normalize_site_constraint(token)
        else:
            cleaned=re.sub(r'[()\[\]{}]+',' ',token)
            if cleaned.strip().lower() in {'and','or','not'} or _is_search_filler_word(cleaned):
                continue
        cleaned=re.sub(r'\s+',' ',cleaned).strip()
        if cleaned and cleaned.lower() not in {'and','or','not'}:
            _append_site_constraint_once(out, cleaned, seen_site)
    return re.sub(r'\s+',' ',' '.join(out)).strip()

def sanitize_local_search_engine_query(value: str) -> str:
    """Constrain Local GPU/search-engine requests to broad, useful syntax.

    Search providers are sensitive to over-constrained exact phrases. Across one
    outgoing request, quote at most three meaningful words total. A quoted phrase
    longer than three words is emitted as normal terms instead. Boolean/filler
    operators (AND/OR/NOT, including and/or forms) are removed completely.
    """
    text=_repair_unbalanced_generated_quotes(strip_external_exclusions(str(value or ''))).replace('|',' ')
    text=re.sub(r'(?i)\b(?:and\s*/\s*or|and/or|and\\or)\b',' ',text)
    tokens=re.findall(r'"[^"\n]*"|site:[^\s]+|[^\s]+',text,flags=re.I)
    out=[]; quoted_words=0; seen_site=[False]
    boolean={'and','or','not','&','&&','||'}
    for token in tokens:
        if token.lower().startswith('site:'):
            cleaned=_normalize_site_constraint(token)
            _append_site_constraint_once(out, cleaned, seen_site)
            continue
        quoted=token.startswith('"') and token.endswith('"')
        raw=token[1:-1] if quoted else token
        raw=re.sub(r'[/\\]+',' ',raw)
        raw=re.sub(r'[()\[\]{}]+',' ',raw)
        words=[w for w in re.split(r'\s+',raw.strip()) if w and w.casefold() not in boolean]
        if not quoted:
            words=[w for w in words if not _is_search_filler_word(w)]
        if not words:
            continue
        # Keep exact matching only while the whole request remains at or below three
        # quoted words. This prevents combinations such as three separate 3-word
        # phrases from becoming a nine-word exact-match search.
        remaining=max(0,3-quoted_words)
        if quoted and len(words)<=remaining and len(words)<=3:
            out.append('"'+' '.join(words)+'"'); quoted_words+=len(words)
        else:
            out.extend(words)
    cleaned=re.sub(r'\s+',' ',' '.join(out)).strip()
    # Defensive final pass: boolean/filler terms must never leave ScoutBox as search
    # operators or standalone words.
    cleaned=re.sub(r'(?i)(?<![A-Za-z0-9])(?:and|or|not)(?![A-Za-z0-9])',' ',cleaned)
    cleaned=re.sub(r'\s+[/\\]+\s+',' ',cleaned)
    return re.sub(r'\s+',' ',cleaned).strip()



CV_TECH_LEXICON = [
    '.NET', '.NET Framework', '.NET Core', 'ASP.NET', 'C#', 'F#', 'VB.NET',
    'WPF', 'WinForms', 'MAUI', 'UWP', 'WCF', 'JavaFX', 'Spring', 'Hibernate',
    'Node.js', 'ReactJS', 'React Native', 'Expo', 'Swift', 'Kotlin', 'PHP', 'HTML/CSS',
    'Bash', 'PowerShell', 'Wix', 'Squarespace', 'UX/UI', 'responsive design',
    'REST', 'SOAP', 'GraphQL', 'Tomcat', 'Nginx', 'Apache',
    'VoIP', 'Asterisk', 'FreePBX', 'SIP', 'PBX', 'RTP', 'RTCP', 'WebRTC',
    'Kamailio', 'OpenSIPS', 'Twilio', 'SMPP', 'SMS', 'GSM', 'LTE', '5G', '3CX',
    'FreeSWITCH', 'RTSP', 'ONVIF', 'CCTV',
    'QEMU', 'KVM', 'Xen', 'VirtualBox', 'VMware', 'Hyper-V', 'virtualization',
    'Firecracker', 'gVisor', 'Kata Containers', 'LXC', 'TAP/TUN',
    'DOSBox', 'DOSBox-X', 'ScummVM', 'Bochs', 'emulation', 'emulator', 'UEFI', 'BIOS',
    'C', 'C++', 'Rust', 'Python', 'JavaScript', 'TypeScript', 'Go', 'Java',
    'Linux', 'Debian', 'Ubuntu', 'Red Hat', 'CentOS', 'Windows', 'Windows Server',
    'FreeBSD', 'OpenWrt', 'Yocto', 'Buildroot', 'Raspberry Pi',
    'STM32', 'ESP32', 'PIC', 'dsPIC', 'AVR', 'ARM', 'Cortex-M', 'RISC-V',
    'FreeRTOS', 'Zephyr', 'RTOS', 'JTAG', 'SWD', 'UART', 'SPI', 'I2C', 'USB',
    'CAN', 'CANopen', 'BACnet', 'BACnet/IP', 'Modbus', 'KNX', 'OPC-UA', 'DNP3', 'BLE',
    'Bluetooth', 'Zigbee', 'LoRa', 'LoRaWAN', 'BMS', 'BAS', 'building automation',
    'building management systems', 'industrial automation', 'SCADA', 'DDC', 'IEC 61131-3', 'Codesys', 'Niagara 4',
    'FPGA', 'Verilog', 'SystemVerilog', 'VHDL', 'RTL', 'ASIC', 'KiCad', 'Altium',
    'PCB', 'schematic', 'signal integrity', 'reverse engineering', 'binary analysis',
    'firmware analysis', 'protocol analysis', 'assembly', 'x86', 'MIPS', 'PowerPC',
    'Qt', 'OpenGL', 'DirectX', 'Vulkan', 'FFmpeg', 'GStreamer', 'SQL', 'PostgreSQL',
    'SQL Server', 'Oracle', 'MySQL', 'MongoDB', 'Redis', 'RabbitMQ', 'MQTT',
    'Django', 'Flask', 'FastAPI', 'React', 'Angular', 'Vue', 'Docker', 'Kubernetes',
    'Ansible', 'Terraform', 'GitLab CI', 'GitHub', 'Bitbucket', 'CI/CD', 'GitOps',
    'GCP', 'Azure', 'AWS', 'Azure Data Factory', 'AWS Glue', 'Microsoft Fabric',
    'MuleSoft', 'Salesforce', 'Ellucian', 'SaaS', 'ETL', 'SSIS', 'dbt', 'Power BI', 'Tableau', 'Cognos',
    'TensorFlow', 'PyTorch', 'CUDA', 'OpenCV', 'scikit-learn', 'NumPy', 'Pandas',
    'Keras', 'MLflow', 'Apache Spark', 'Jupyter', 'Prometheus', 'Grafana',
    'PayPal', 'Stripe', 'Worldpay', '3-D Secure', '3DS', 'Jira', 'Confluence', 'Scrum', 'Agile',
    'SMTP', 'IMAP',
]


_PROFILE_CONCEPT_CONSTRAINT_RE = re.compile(
    r'(?i)\b(?:avoid|exclude|excluding|prefer|preference|preferred|salary|pay|rate|'
    r'low[- ]?rate|marketplace|funnel|interview|leetcode|psychometric|onsite only|'
    r'remote only|full[- ]?time only|part[- ]?time only|company size|hiring process)\b'
)


def _profile_concept_supported(term: str, corpus: str) -> bool:
    """Require Candidate Profile concepts to be Resume-grounded technical/domain terms.

    Older extraction accepted a concept when *one* token occurred anywhere in a CV. That
    allowed preference prose such as "long enterprise funnels and low-rate marketplaces"
    to leak into Resume concepts because the word "enterprise" happened to be present.
    """
    term=_normalise(term).strip(' .,:;-')
    if not term or _PROFILE_CONCEPT_CONSTRAINT_RE.search(term):
        return False
    low_corpus=(corpus or '').casefold()
    low_term=term.casefold()
    if low_term in low_corpus:
        return True
    technical_known={x.casefold() for x in CV_TECH_LEXICON}
    technical_known.update(x.casefold() for x in SKILL_GRAPH.keys())
    if low_term in technical_known:
        return any(_has(corpus, x) for x in [term])
    generic={
        'and','with','from','using','used','work','working','system','systems','software',
        'engineering','engineer','technical','development','developer','technology','technologies',
        'platform','platforms','solution','solutions','enterprise','production','project','projects',
    }
    tokens=[x.casefold() for x in re.findall(r'[A-Za-z0-9+#./-]{3,}',term) if x.casefold() not in generic]
    if not tokens:
        return False
    matched=sum(1 for token in tokens if token in low_corpus)
    return matched >= max(1, (len(tokens)+1)//2)

def _resume_specific_terms(cv_texts: list[str]) -> list[dict]:
    """Extract useful technical terms that are present in CV text but absent from SKILL_GRAPH.

    This is intentionally content-driven. It captures punctuation-heavy platforms/languages,
    CV-listed technologies and recurrent technical tokens instead of depending on a fixed
    role-name dictionary.  The returned list is intentionally broad because search planning
    rotates across it over time.
    """
    counts=Counter()
    display={}
    lexicon_keys={x.lower(): x for x in CV_TECH_LEXICON}
    stop={'http','https','www','email','phone','experience','years','skills','technical','engineering','engineer','software','system','systems','project','projects','work','working','using','used','development','developed','support','based','including','with','from','and','the','for','this','that','resume','profile','client','company'}

    def remember(term, seen):
        term=_normalise(term).strip(' ,.;:()[]{}')
        low=term.lower()
        if not term or low in stop or len(term)<2 or len(term)>60 or low in seen:
            return
        seen.add(low); counts[low]+=1; display.setdefault(low,lexicon_keys.get(low,term))

    for text in cv_texts:
        text=text or ''
        seen=set()
        # Explicit CV technology lexicon catches normal mixed-case words such as Asterisk
        # and punctuation-heavy terms such as .NET/C# that generic tokenization often misses.
        for term in CV_TECH_LEXICON:
            if _has(text,term):
                remember(term,seen)
        # Also keep unusual tokens/acronyms/platform names that appear in the CV but are not
        # already modelled in SKILL_GRAPH.
        candidates=re.findall(r'(?<!\w)(?:(?i:\.NET(?:\s+(?:Framework|Core))?)|\.?[A-Za-z][A-Za-z0-9]*(?:[+#./-][A-Za-z0-9+#.-]+)+|[A-Z][A-Z0-9]{1,9}|[A-Za-z]+\s+(?:Framework|SDK|BIOS|UEFI|Assembly|Assembler|Runtime|Server|Database|Protocol)|(?:Assembly|Assembler|UEFI))(?!\w)',text)
        for raw in candidates:
            remember(raw,seen)

    known=' '.join(SKILL_GRAPH.keys()).lower()
    rows=[]
    for low,count in counts.most_common(100):
        if low in known or count<1:
            continue
        # CV lexicon hits are strong even when they appear only once; the planner rotates
        # them rather than over-weighting every run.
        score=8.0+min(9,count*1.25)
        rows.append({'term':display[low],'score':round(score,2),'category':'cv-specific','aliases':[],'matched_variants':[display[low]],'cv_documents':count,'sources':{'cv':True,'high_preference':False,'medium_preference':False,'low_preference':False,'campaign':False}})
    return rows[:80]


def grounded_resume_concepts(cv_texts: list[str], limit=PROFILE_CONCEPT_LIMIT) -> list[str]:
    """Return a broad, deterministic, Resume-only technical/domain vocabulary.

    This deliberately excludes Candidate Profile preference prose and campaign steering.
    It is used when rebuilding the editable Candidate Profile so core CV technologies are
    retained even if the AI summarizer omits them.
    """
    rows=[]
    for canonical,spec in SKILL_GRAPH.items():
        variants=[canonical]+list(spec.get('patterns') or [])+list(spec.get('aliases') or [])
        doc_hits=sum(1 for text in cv_texts if any(_has(text,v) for v in variants))
        if doc_hits:
            rows.append((float(spec.get('weight') or 0)+doc_hits*3.0,canonical))

    # Keep literal Resume technologies as first-class concepts even when a broader graph
    # concept already covers them. For example, a CV mentioning BACnet and Modbus should
    # retain both literal protocol names as well as the broader BACnet / Modbus concept.
    # This makes the editable profile useful for exact-keyword discovery and diagnostics.
    for term in CV_TECH_LEXICON:
        doc_hits=sum(1 for text in cv_texts if _has(text,term))
        if doc_hits:
            rows.append((10.0+doc_hits*2.0,term))

    for row in _resume_specific_terms(cv_texts):
        rows.append((float(row.get('score') or 0),str(row.get('term') or '')))
    out=[]; seen=set()
    for _score,term in sorted(rows,key=lambda x:(-x[0],str(x[1]).casefold())):
        term=_normalise(term).strip().lower()
        if not term or term in seen:
            continue
        seen.add(term); out.append(term)
        if len(out)>=max(1,int(limit or PROFILE_CONCEPT_LIMIT)):
            break
    return out


def grounded_resume_roles(cv_texts: list[str], limit=PROFILE_ROLE_LIMIT) -> list[str]:
    """Infer built-in role families from Resume evidence only (no preferences/campaign)."""
    present={x.casefold() for x in grounded_resume_concepts(cv_texts,PROFILE_CONCEPT_LIMIT)}
    scored=[]
    for role,triggers in ROLE_FAMILIES.items():
        hits=sum(1 for trigger in triggers if trigger.casefold() in present)
        if hits:
            scored.append((hits,role))
    return [role for _hits,role in sorted(scored,key=lambda x:(-x[0],x[1].casefold()))[:max(1,int(limit or PROFILE_ROLE_LIMIT))]]

def _campaign_negative_terms(campaign) -> list[str]:
    if not campaign:
        return []
    out=[]
    for raw in re.split(r'[,;\n|]+',str(getattr(campaign,'negative_constraints','') or '')):
        term=raw.strip().lstrip('-').strip(' "\'')
        if not term or _GENERIC_FULL_STACK_RE.search(term):
            continue
        if term.lower() not in {x.lower() for x in out}:
            out.append(term)
    return out[:40]

def build_search_profile(campaign=None, use_saved_overrides=True) -> dict:
    """Build a deterministic search profile from all active Resumes plus user preferences.

    Campaign role families/technologies are steering inputs.  Resume evidence remains the
    dominant score source so an over-broad campaign cannot drown out the candidate's
    actual specialist background.
    """
    profile = Profile.objects.get_or_create(pk=1)[0]
    docs = extract_active_cv_texts()
    cv_texts = [d['text'] for d in docs]
    high = profile.high_priority_text or ''
    medium = profile.medium_priority_text or ''
    low = profile.low_priority_text or ''
    campaign_roles = _split_config(getattr(campaign, 'role_families', ''))
    campaign_tech = _split_config(getattr(campaign, 'technologies', ''))
    campaign_blob = ' '.join(campaign_roles + campaign_tech)

    skills = []
    for canonical, spec in SKILL_GRAPH.items():
        variants = [canonical] + spec['patterns'] + spec['aliases']
        doc_hits = 0
        occurrence_count = 0
        matched_variants = []
        for text in cv_texts:
            matched_doc = False
            for variant in variants:
                if _has(text, variant):
                    occurrence_count += 1
                    matched_doc = True
                    if variant not in matched_variants:
                        matched_variants.append(variant)
            if matched_doc:
                doc_hits += 1

        pref_high = any(_has(high, x) for x in variants)
        pref_med = any(_has(medium, x) for x in variants)
        pref_low = any(_has(low, x) for x in variants)
        campaign_match = any(_has(campaign_blob, x) or _has(x, campaign_blob) for x in variants if campaign_blob)
        if not (doc_hits or pref_high or pref_med or pref_low or campaign_match):
            continue

        # Resume evidence dominates.  Cross-CV recurrence is useful because multiple tailored
        # CVs repeatedly mentioning a skill is strong evidence of a genuine specialty.
        score = 0.0
        if doc_hits:
            score += spec['weight'] + doc_hits * 3.2 + min(5.0, occurrence_count * 0.6)
        if pref_high:
            score += 8.0
        if pref_med:
            score += 4.0
        if pref_low:
            score += 1.5
        if campaign_match:
            score += 3.0
        skills.append({
            'term': canonical,
            'score': round(score, 2),
            'category': spec['category'],
            'aliases': list(dict.fromkeys(spec['aliases'])),
            'matched_variants': matched_variants[:12],
            'cv_documents': doc_hits,
            'sources': {
                'cv': bool(doc_hits), 'high_preference': pref_high, 'medium_preference': pref_med,
                'low_preference': pref_low, 'campaign': campaign_match,
            },
        })

    # Preserve specialist technologies found directly in the uploaded CVs even when the
    # built-in graph does not know them (for example .NET Framework or unusual SDKs).
    existing_terms={x['term'].lower() for x in skills}
    for row in _resume_specific_terms(cv_texts):
        if row['term'].lower() not in existing_terms:
            skills.append(row); existing_terms.add(row['term'].lower())

    # Short explicit free-text preferences that are not represented in the built-in graph
    # also steer search. Resume-derived graph skills still receive substantially larger scores.
    known_terms={x['term'].lower() for x in skills}
    for raw,boost,source_name in [
        *[(x,7.0,'high_preference') for x in _explicit_preference_phrases(high)],
        *[(x,4.0,'medium_preference') for x in _explicit_preference_phrases(medium)],
        *[(x,1.5,'low_preference') for x in _explicit_preference_phrases(low)],
    ]:
        if raw.lower() in known_terms: continue
        skills.append({'term':raw,'score':boost,'category':'preference','aliases':[],'matched_variants':[], 'cv_documents':0,
                       'sources':{'cv':False,'high_preference':source_name=='high_preference','medium_preference':source_name=='medium_preference','low_preference':source_name=='low_preference','campaign':False}})
        known_terms.add(raw.lower())

    # Explicit campaign terms that are outside the built-in graph are retained as weak
    # steering signals. They never outrank strong Resume terms by themselves.
    known_blob = ' '.join(x['term'] for x in skills).lower()
    for raw in campaign_tech:
        if raw.lower() not in known_blob:
            skills.append({'term': raw, 'score': 5.0, 'category': 'campaign', 'aliases': [], 'matched_variants': [], 'cv_documents': 0,
                           'sources': {'cv': False, 'high_preference': False, 'medium_preference': False, 'low_preference': False, 'campaign': True}})

    skills.sort(key=lambda x: (-x['score'], x['term'].lower()))
    score_by_term = {x['term']: x['score'] for x in skills}

    inferred_roles = []
    for role, triggers in ROLE_FAMILIES.items():
        vals = [score_by_term[t] for t in triggers if t in score_by_term]
        if vals:
            inferred_roles.append({'role': role, 'score': round(sum(sorted(vals, reverse=True)[:3]) / min(3, len(vals)), 2), 'source': 'cv'})
    for role in campaign_roles:
        # Configured role families steer the planner but do not replace inferred CV families.
        existing = next((x for x in inferred_roles if x['role'].lower() == role.lower()), None)
        if existing:
            existing['score'] += 2.5
            existing['source'] = 'cv+campaign'
        else:
            inferred_roles.append({'role': role, 'score': 6.0, 'source': 'campaign'})
    inferred_roles.sort(key=lambda x: (-x['score'], x['role'].lower()))

    # Candidate Profile values are additive/curated signals, not a whitelist that erases
    # richer deterministic Resume evidence. This keeps manual/profile edits useful while
    # ensuring discovery still sees technologies and role families present in the source CV.
    saved_scope=profile.scope_json or {}
    if use_saved_overrides and 'cv_concepts' in saved_scope:
        requested=[]; seen=set()
        for raw in saved_scope.get('cv_concepts') or []:
            term=_normalise(str(raw)).lower()
            if term and term not in seen:
                seen.add(term); requested.append(term)
        existing={x['term'].lower():x for x in skills}
        for term in requested[:PROFILE_CONCEPT_LIMIT]:
            if term in existing:
                row=existing[term]
                row['sources']=dict(row.get('sources') or {}, profile_override=True)
                row['score']=max(float(row.get('score') or 0),8.0)+1.0
            else:
                row={'term':term,'score':8.0,'category':'profile','aliases':[],'matched_variants':[],'cv_documents':0,'sources':{'cv':False,'high_preference':False,'medium_preference':False,'low_preference':False,'campaign':False,'profile_override':True}}
                skills.append(row); existing[term]=row
        skills.sort(key=lambda x:(-x['score'],x['term'].lower()))
    if use_saved_overrides and 'likely_roles' in saved_scope:
        requested=[]; seen=set()
        for raw in saved_scope.get('likely_roles') or []:
            role=_normalise(str(raw)).lower()
            if role and role not in seen:
                seen.add(role); requested.append(role)
        existing={x['role'].lower():x for x in inferred_roles}
        for role in requested[:PROFILE_ROLE_LIMIT]:
            if role in existing:
                row=existing[role]
                row['source']='cv+profile' if str(row.get('source') or '').startswith('cv') else 'profile'
                row['score']=max(float(row.get('score') or 0),8.0)+1.0
            else:
                row={'role':role,'score':8.0,'source':'profile'}
                inferred_roles.append(row); existing[role]=row
        inferred_roles.sort(key=lambda x:(-x['score'],x['role'].lower()))

    # Preserve diversity across multiple role-tailored Resumes. Repeated common skills are
    # valuable, but they should not completely drown out a Technical-Writing or Retro
    # CV that exists specifically to represent another opportunity family.
    cv_focus=[]
    for doc in docs:
        label=(doc.get('label') or doc.get('name') or '').lower()
        focus_scores=[]
        for canonical,spec in SKILL_GRAPH.items():
            variants=[canonical]+spec['patterns']+spec['aliases']
            if not any(_has(doc['text'],v) for v in variants):
                continue
            fscore=float(spec['weight'])
            for hint,hint_terms in LABEL_FOCUS_HINTS.items():
                if hint in label and canonical in hint_terms:
                    fscore += 15
            focus_scores.append((fscore,canonical,spec['category']))
        focus_scores.sort(reverse=True)
        chosen=[]; cats=set()
        for fscore,term,cat in focus_scores:
            # Prefer some category diversity within one CV, but allow the top specialty
            # to contribute two closely related concepts where useful.
            if len(chosen)>=4: break
            if cat in cats and len(chosen)>=2: continue
            chosen.append({'term':term,'score':round(fscore,2),'category':cat}); cats.add(cat)
        if chosen:
            cv_focus.append({'id':doc['id'],'label':doc['label'],'terms':chosen})

    return {
        'cv_count': len(docs),
        'cv_labels': [d['label'] for d in docs],
        'cv_focus': cv_focus,
        'skills': skills[:PROFILE_CONCEPT_LIMIT],
        'role_families': inferred_roles[:PROFILE_ROLE_LIMIT],
        'negative_terms': list(dict.fromkeys(_preference_negatives(profile) + _campaign_negative_terms(campaign))),
        'operating_location': profile.operating_location,
        'preference_excerpt': {
            'high': high[:600], 'medium': medium[:600], 'low': low[:600],
        },
        'campaign_steering': {'roles': campaign_roles, 'technologies': campaign_tech},
    }


def _quote(term: str) -> str:
    term = term.strip()
    if not term:
        return ''
    return f'"{term}"' if ' ' in term or '/' in term or '+' in term else term


def _stable_rotation(items: list[dict], campaign_id: int | None, interval_minutes: int, rotation_offset: int = 0) -> list[dict]:
    if not items:
        return []
    interval = max(30, int(interval_minutes or 120))
    bucket = int(time.time() // (interval * 60)) + int(rotation_offset or 0)
    seed = f'{campaign_id or 0}:{bucket}'
    return sorted(items, key=lambda x: hashlib.sha256((seed + '|' + x['query']).encode('utf-8')).hexdigest())


def _rotation_bucket(interval_minutes: int, rotation_offset: int = 0) -> int:
    interval=max(30,int(interval_minutes or 120))
    return int(time.time()//(interval*60))+int(rotation_offset or 0)


def _coverage_rotation(items: list[dict], campaign_id: int | None, interval_minutes: int, slots: int, rotation_offset: int = 0) -> list[dict]:
    """Rotate a source list with a hard bounded-coverage guarantee.

    Hash sorting is useful for variety but cannot guarantee that the last source is ever
    selected when only the first N rows are taken.  Source coverage instead advances by
    exactly the number of source slots per interval, so every source is selected within
    ``ceil(len(items) / slots)`` consecutive rotations.
    """
    if not items:
        return []
    slots=max(1,min(len(items),int(slots or 1)))
    bucket=_rotation_bucket(interval_minutes,rotation_offset)
    start=((int(campaign_id or 0)%len(items)) + (bucket*slots)) % len(items)
    return items[start:]+items[:start]


def _source_target_rows(target_domains=None, target_sources=None, *, campaign_id=None, interval_minutes=120, rotation_offset=0, per_source_limit=6) -> list[dict]:
    """Return one rotating domain variant for every enabled source identity.

    Expansion happens only *after* the source identity has survived scheduling.  This
    prevents a multi-domain source such as Indeed/Lever from consuming several slots
    before later sources (for example YC Work at a Startup) get a single chance.
    """
    raw=[]
    if target_sources:
        for item in target_sources:
            if isinstance(item,dict):
                raw.append({
                    'source_id':item.get('source_id'),
                    'source_name':str(item.get('source_name') or item.get('name') or '').strip(),
                    'domain':str(item.get('domain') or item.get('base_domain') or '').strip(),
                })
            else:
                raw.append({'source_id':None,'source_name':'','domain':str(item or '').strip()})
    else:
        raw=[{'source_id':None,'source_name':'','domain':str(value or '').strip()} for value in (target_domains or [])]

    normalized=[]; seen=set()
    for item in raw:
        domain=_normalise(str(item.get('domain') or '')).lower().removeprefix('www.')
        if '://' in domain:
            try:
                from urllib.parse import urlsplit
                domain=(urlsplit(domain).hostname or '').lower().removeprefix('www.')
            except Exception:
                domain=''
        domain=domain.strip('/').split('/',1)[0].split(':',1)[0]
        if not domain:
            continue
        identity=('id',str(item.get('source_id'))) if item.get('source_id') not in (None,'') else ('domain',domain)
        if identity in seen:
            continue
        seen.add(identity)
        normalized.append({**item,'domain':domain})

    bucket=_rotation_bucket(interval_minutes,rotation_offset)
    out=[]
    for idx,item in enumerate(normalized):
        variants=expand_source_domains(item['domain'],limit=max(1,int(per_source_limit or 6))) or [item['domain']]
        identity=f"{item.get('source_id') or ''}|{item.get('source_name') or ''}|{item['domain']}"
        salt=int(hashlib.sha256(identity.encode('utf-8')).hexdigest()[:8],16)
        domain=variants[(bucket+salt)%len(variants)]
        out.append({
            'source_id':item.get('source_id'),
            'source_name':item.get('source_name') or item['domain'],
            'source_domain':item['domain'],
            'domain':domain,
        })
    return out


def _compatible_pairs(skills: list[dict]) -> list[tuple[dict, dict]]:
    pairs = []
    for i, a in enumerate(skills[:16]):
        for b in skills[i + 1:16]:
            # Same-category pairs find specific niches; selected cross-category pairs expose
            # valuable intersections such as QEMU+reverse-engineering or embedded+protocols.
            same = a['category'] == b['category'] or a['category'].split('-')[0] == b['category'].split('-')[0]
            cross = {a['category'].split('-')[0], b['category'].split('-')[0]} in [
                {'virtualization', 'reverse'}, {'retro', 'reverse'}, {'embedded', 'reverse'},
                {'embedded', 'systems'}, {'retro', 'virtualization'}, {'writing', 'embedded'},
                {'writing', 'reverse'},
            ]
            if same or cross:
                pairs.append((a, b))
    return pairs


def build_query_plan(campaign=None, max_queries=12, target_domains: Iterable[str] | None = None, interval_minutes=120, rotation_offset=0) -> dict:
    """Return a Resume-first rotating query plan and explanatory metadata."""
    search_profile = build_search_profile(campaign)
    skills = search_profile['skills']
    roles = search_profile['role_families']
    candidates = []

    def add(query, terms, kind, score, rationale):
        query = sanitize_local_search_engine_query(query)
        if not query:
            return
        candidates.append({'query': query, 'terms': terms[:3], 'kind': kind, 'score': round(score, 2), 'rationale': rationale})

    top_skills = skills[:20]
    top_roles = roles[:8]
    if not top_skills and not top_roles:
        return {
            'queries':[], 'search_profile':search_profile,
            'location_policy':{'operating_location':search_profile.get('operating_location',''),'location_used_in_queries':False,'note':'Candidate Profile contains no reliable discovery vocabulary; review Resume concepts / likely roles.'},
            'profile_review_required':True,
        }

    # Give each role-tailored Resume a chance to contribute its own search family. These
    # candidates are rotated so six CVs do not consume every slot in every run.
    for focus in search_profile.get('cv_focus',[]):
        choices=focus.get('terms',[])
        if not choices: continue
        primary=choices[0]
        secondary=next((x for x in choices[1:] if primary['term'].lower() not in x['term'].lower() and x['term'].lower() not in primary['term'].lower()), choices[1] if len(choices)>1 else None)
        chosen=[primary]+([secondary] if secondary else [])
        terms=[x['term'] for x in chosen]
        query=' '.join(_quote(t) for t in terms)+' remote'
        add(query,terms,'cv-document-focus',sum(x['score'] for x in chosen)+8,
            f'Focus preserved from Resume variant: {focus.get("label","Resume")}.')

    # When a CV yields concepts but no defensible role family, search the evidenced
    # concepts directly rather than inventing a software/embedded role.
    if not top_roles:
        for skill in top_skills[:10]:
            add(f'{_quote(skill["term"])} hiring remote', [skill['term']], 'skill-only', skill['score'], 'Resume/Profile concept without an inferred role; no profession fallback was invented.')

    # Role + one strong Resume term.  This is the backbone of normal vacancy discovery.
    for role in top_roles[:6]:
        for skill in top_skills[:8]:
            if role['source'] == 'campaign' and not skill.get('sources', {}).get('cv', True):
                continue
            add(f'{_quote(role["role"])} {_quote(skill["term"])} remote', [role['role'], skill['term']], 'role-skill',
                role['score'] + skill['score'], f'Resume/inferred role family + strong Resume skill ({skill["term"]}).')

    # Explicit role-family / technology controls are steering inputs, not a replacement
    # keyword list. Reserve candidates that pair one configured term with one strong CV
    # concept so user-selected PIC/STM32/etc. actually receives coverage without dominating.
    steering=search_profile.get('campaign_steering') or {}
    for raw in (steering.get('technologies') or [])[:8]:
        related=None
        for skill in top_skills:
            variants=[skill['term']]+skill.get('aliases',[])+skill.get('matched_variants',[])
            if any(_has(raw,v) or _has(v,raw) for v in variants):
                related=skill; break
        related=related or top_skills[0]
        companion=related['term']
        if _has(companion,raw) or _has(raw,companion):
            basecat=related.get('category','').split('-')[0]
            alt=next((x['term'] for x in top_skills if x['term']!=related['term'] and x.get('category','').split('-')[0]==basecat and not (_has(x['term'],raw) or _has(raw,x['term']))),None)
            companion=alt or next((a for a in related.get('aliases',[]) if not (_has(a,raw) or _has(raw,a))), top_skills[1]['term'] if len(top_skills)>1 else 'systems')
        add(f'{_quote(raw)} {_quote(companion)} remote', [raw,companion], 'campaign-steering',
            related['score']+7, f'Configured technology {raw} woven into Resume-derived concept {companion}.')
    for raw in (steering.get('roles') or [])[:6]:
        related=top_skills[0]
        add(f'{_quote(raw)} {_quote(related["term"])} remote', [raw,related['term']], 'campaign-steering',
            related['score']+5, f'Configured role family {raw} woven into Resume-derived search.')

    # Niche intersections using two CV signals, not a large AND-list.
    for a, b in _compatible_pairs(top_skills):
        add(f'{_quote(a["term"])} {_quote(b["term"])} remote', [a['term'], b['term']], 'niche-intersection',
            a['score'] + b['score'] + 2, f'Rare Resume intersection: {a["term"]} + {b["term"]}.')

    # Synonym/adjacency expansion casts a wider net. Prefer variants that were not literally
    # observed in the uploaded Resumes, so searches are not just duplicates of Resume wording.
    for skill in top_skills[:18]:
        matched = {x.lower() for x in skill.get('matched_variants', [])}
        aliases = [a for a in skill.get('aliases', []) if a.lower() not in matched] or skill.get('aliases', [])
        for alias in aliases[:3]:
            add(f'{_quote(alias)} {_quote(skill["term"])} remote', [skill['term'], alias], 'synonym-expansion',
                skill['score'] + 3, f'Wider-net wording for Resume skill {skill["term"]}: {alias}.')

    # Problem-signal searches find consulting/help/maintainer needs that are never posted as jobs.
    problem_phrases = ['looking for someone', 'need help with', 'looking for a contractor', 'consultant', 'maintainer wanted']
    for skill in top_skills[:16]:
        for phrase in problem_phrases[:3]:
            add(f'"{phrase}" {_quote(skill["term"])}', [phrase, skill['term']], 'problem-signal',
                skill['score'] + 1.5, f'Hidden Leads/problem-language search around {skill["term"]}.')

    # Contract/project variants matter for specialist work and do not require geography in query.
    for skill in top_skills[:16]:
        add(f'{_quote(skill["term"])} remote contract', [skill['term'], 'contract'], 'engagement',
            skill['score'] + 1, f'Specialist contract variant for {skill["term"]}.')

    # Target selected source domains using the strongest Resume-derived base queries.
    # Schedule source identities first, then choose one rotating domain variant for each.
    # This legacy planner is currently not the Local AI path, but it follows the same
    # bounded-coverage rule so future callers cannot reintroduce pre-rotation starvation.
    base_snapshot = sorted(candidates, key=lambda x: -x['score'])[:12]
    source_targets=_source_target_rows(target_domains=target_domains,campaign_id=getattr(campaign,'pk',None),interval_minutes=interval_minutes,rotation_offset=rotation_offset)
    source_targets=_coverage_rotation(source_targets,getattr(campaign,'pk',None),interval_minutes,min(10,len(source_targets) or 1),rotation_offset)[:10]
    for target in source_targets:
        domain=target['domain']
        for item in base_snapshot[:3]:
            add(f'site:{domain} {item["query"]}', item['terms'], 'site-targeted', item['score'] + 0.5,
                f'Resume-first query targeted at enabled source {target.get("source_name") or domain} via {domain}.')

    # Negative preferences are intentionally enforced after retrieval. Search-engine
    # exclusion syntax can suppress useful niche results and differs by provider.

    # Optional advanced campaign text remains supported, but is not required and is deliberately
    # lower priority than Resume-derived queries.
    extra = (getattr(campaign, 'extra_text', '') if campaign else '').strip()
    if extra:
        add(extra, [extra[:80]], 'advanced-manual', 3.0, 'Optional advanced campaign text.')

    # De-duplicate, then deliberately reserve room for wider-net variants. A planner that
    # always chooses only its highest-scoring literal Resume terms would never exercise the
    # synonym graph, so each run mixes core intersections, synonym expansion and hidden-
    # market/problem-language queries. Rotation changes which members of each bucket run.
    unique = {}
    for item in candidates:
        key = item['query'].lower()
        if key not in unique or item['score'] > unique[key]['score']:
            unique[key] = item
    ranked = sorted(unique.values(), key=lambda x: -x['score'])

    cv_bucket=[x for x in ranked if x['kind']=='cv-document-focus']
    core=[x for x in ranked if x['kind'] in ('role-skill','niche-intersection')]
    steering_bucket=[x for x in ranked if x['kind']=='campaign-steering']
    expansions=[x for x in ranked if x['kind']=='synonym-expansion']
    hidden=[x for x in ranked if x['kind'] in ('problem-signal','engagement')]
    targeted=[x for x in ranked if x['kind']=='site-targeted']
    manual=[x for x in ranked if x['kind']=='advanced-manual']
    selected=[]
    cv_take=min(len(cv_bucket),max(1,round(max_queries*0.25))) if cv_bucket else 0
    steering_take=min(1,len(steering_bucket)) if steering_bucket else 0
    core_take=max(1,round(max_queries*0.20))
    expansion_take=max(1,round(max_queries*0.25))
    hidden_take=max(1,round(max_queries*0.17))
    targeted_take=max(0,max_queries-(cv_take+steering_take+core_take+expansion_take+hidden_take))
    allocations=[
        (cv_bucket,cv_take),
        (steering_bucket,steering_take),
        (core,core_take),
        (expansions,expansion_take),
        (hidden,hidden_take),
        (targeted,targeted_take),
    ]
    for bucket,take in allocations:
        ordered=_stable_rotation(bucket,getattr(campaign,'pk',None),interval_minutes,rotation_offset)
        selected.extend(ordered[:max(0,min(take,max_queries-len(selected)))])
        if len(selected)>=max_queries: break
    # Optional manual context is allowed but cannot displace the Resume-first mix.
    if manual and len(selected)<max_queries:
        selected.append(manual[0])
    if len(selected)<max_queries:
        for item in _stable_rotation(ranked,getattr(campaign,'pk',None),interval_minutes,rotation_offset):
            if item not in selected:
                selected.append(item)
            if len(selected)>=max_queries: break

    return {
        'queries': selected[:max_queries],
        'search_profile': search_profile,
        'candidate_query_count': len(ranked),
        'location_used_in_queries': False,
        'location_policy': 'Operating location is used only for eligibility/fit scoring after discovery.',
        'query_policy': 'Resume-first; 1-3 specialist concepts per query; role/technology settings steer rather than replace Resume evidence; synonyms/adjacencies rotate across runs.',
        'rotation_offset': int(rotation_offset or 0),
        'rotation_interval_minutes': max(30, int(interval_minutes or 120)),
        'rotation_bucket': int(time.time() // (max(30, int(interval_minutes or 120)) * 60)) + int(rotation_offset or 0),
    }



ROLE_ALIAS_FALLBACKS = {
    'technical writer': ['technical writer','technical author','documentation engineer','developer documentation writer','technical content writer','documentation specialist'],
    'technical writing': ['technical writer','technical author','documentation engineer','developer documentation writer','technical content writer'],
    'embedded software engineer': ['embedded software engineer','embedded systems engineer','firmware engineer','embedded developer','embedded C/C++ engineer'],
    'firmware engineer': ['firmware engineer','embedded firmware engineer','embedded software engineer','microcontroller engineer','firmware developer'],
    'reverse engineer': ['reverse engineer','reverse engineering specialist','binary analyst','firmware reverse engineer','security researcher'],
    'systems software engineer': ['systems software engineer','systems programmer','low-level software engineer','platform software engineer'],
    'emulation engineer': ['emulation engineer','emulator developer','virtual hardware engineer','device emulation engineer'],
    'virtualization engineer': ['virtualization engineer','hypervisor engineer','virtualization developer','QEMU engineer'],
    'legacy systems specialist': ['legacy systems specialist','legacy software engineer','modernization engineer','compatibility engineer'],
    'technical trainer': ['technical trainer','technical instructor','developer educator','technical educator'],
}


def _role_alias_fallback(role: str) -> list[str]:
    role=_normalise(role).lower()
    out=list(ROLE_ALIAS_FALLBACKS.get(role,[]))
    if role and role not in out: out.insert(0,role)
    if 'engineer' in role:
        out.extend([role.replace('engineer','developer'),role.replace('engineer','specialist')])
    if 'writer' in role:
        out.extend([role.replace('writer','author'),role.replace('writer','documentation specialist')])
    return list(dict.fromkeys(_normalise(x) for x in out if _normalise(x)))[:8]


def _engagement_context_terms(campaign, profile=None) -> list[str]:
    values=[]
    raw=[]
    for attr in ('engagement_types','company_sizes'):
        value=getattr(campaign,attr,[]) if campaign else []
        raw.extend(value if isinstance(value,list) else _split_config(str(value or '')))
    # Global Engagement Preferences are also query-shaping evidence in Local AI Discovery
    # mode. Campaign-specific choices are kept first and therefore win when both exist.
    scope=dict(getattr(profile,'scope_json',{}) or {}) if profile else {}
    for key in ('engagement','company_size'):
        value=scope.get(key) or []
        raw.extend(value if isinstance(value,list) else _split_config(str(value or '')))
    blob=' '.join(str(x) for x in raw).lower()
    if any(x in blob for x in ('startup','small','1-10','1–10','11-50','11–50','micro')):
        values += ['startup','small team','early stage','new team']
    if any(x in blob for x in ('contract','consult','freelance','project')):
        values += ['contract','contractor','consultant','project work']
    if any(x in blob for x in ('part-time','part time','fractional')):
        values += ['part-time','fractional']
    if any(x in blob for x in ('full-time','full time','permanent')):
        values += ['full-time','hiring']
    # Retain concise operator-selected values even when they are not in the common map.
    for item in raw:
        item=_normalise(str(item)).strip()
        if 1 <= len(item.split()) <= 4 and item.lower() not in {'any','all','no preference'}:
            values.append(item)
    return list(dict.fromkeys(values))[:16]


def _local_query_language(campaign, profile, roles, engagement_terms) -> dict:
    """Optionally ask local Ollama for search-language aliases; never depend on it.

    The local model only proposes wording. It cannot remove the configured campaign role,
    add geography, or change Cloud Web behavior. Results are cached in Profile.scope_json.
    """
    roles=[_normalise(x) for x in roles if _normalise(x)][:8]
    if not roles: return {'role_aliases':{},'context_terms':[]}
    payload={
        'roles':roles,
        'high':(profile.high_priority_text or '')[:700],
        'medium':(profile.medium_priority_text or '')[:700],
        'low':(profile.low_priority_text or '')[:700],
        'engagement':engagement_terms[:16],
    }
    data={'role_aliases':{},'context_terms':[]}
    try:
        # Query planning is an internal helper for URL discovery. It must inherit the
        # exact visible URL Discovery Primary/Fallback route instead of consulting the
        # Ollama provider default or the first model returned by /api/tags.
        from portal.services.ai import route_for_stage, generate_with_route, local_ai_lane_busy
        route=route_for_stage('url_scrape') or {}
        if local_ai_lane_busy():
            return data
        if route.get('provider')!='ollama' or not route.get('model'):
            return data
        route_sig={k:route.get(k,'') for k in ('provider','model','fallback_provider','fallback_model')}
        key=hashlib.sha256(json.dumps({'payload':payload,'route':route_sig},sort_keys=True,ensure_ascii=False).encode()).hexdigest()
        scope=dict(profile.scope_json or {}); cached=scope.get('_local_query_language') or {}
        if cached.get('hash')==key and isinstance(cached.get('data'),dict):
            return cached['data']
        prompt=(
            'Prepare search-engine wording for specialist job discovery. Return JSON only. '
            'For each supplied target role, give 3-6 realistic alternative job titles that employers may use. '
            'Also give up to 10 short search context terms suggested by the priorities and engagement preferences, such as startup, small team, contractor, hiring or documentation. '
            'Do not add locations. Do not replace the target roles with unrelated skills. Do not add exclusions. '
            'Schema: {"role_aliases":{"target role":["alias"]},"context_terms":["term"]}.\n\n'+json.dumps(payload,ensure_ascii=False)
        )
        raw=generate_with_route(route,prompt,stage='query_planning',
                                subject={'type':'campaign','id':getattr(campaign,'pk',''),'label':getattr(campaign,'name','Query planning')},
                                limits_override={'max_input_tokens':1800,'max_output_tokens':600})
        raw=str(raw or '').strip()
        parsed=json.loads(raw)
        aliases={}
        for role in roles:
            vals=(parsed.get('role_aliases') or {}).get(role) or (parsed.get('role_aliases') or {}).get(role.lower()) or []
            if isinstance(vals,str): vals=[vals]
            clean=[]
            for x in vals:
                x=_normalise(str(x)).strip(' .,:;')
                if 1 <= len(x.split()) <= 7 and x.lower() not in {y.lower() for y in clean}: clean.append(x)
            aliases[role]=clean[:6]
        context=[]
        for x in parsed.get('context_terms') or []:
            x=_normalise(str(x)).strip(' .,:;')
            if 1 <= len(x.split()) <= 4 and x.lower() not in {y.lower() for y in context}: context.append(x)
        data={'role_aliases':aliases,'context_terms':context[:10]}
        scope['_local_query_language']={'hash':key,'data':data,'updated_at':time.time()}
        profile.scope_json=scope; profile.save(update_fields=['scope_json','updated_at'])
    except Exception:
        pass
    return data


def build_source_guided_query_plan(campaign=None, max_queries=12, target_domains: Iterable[str] | None = None, custom_domains: Iterable[str] | None = None, target_sources: Iterable[dict] | None = None, interval_minutes=120, rotation_offset=0) -> dict:
    """Build role-anchored queries for Local AI Discovery only.

    Every generated discovery query keeps the campaign's target role (or a realistic role
    alias) as its anchor. CV evidence, all priority tiers, engagement/company-size choices,
    source domains and controlled opportunity-language variants enrich that role instead of
    displacing it. Cloud Web deliberately continues to use build_query_plan().
    """
    max_queries=max(1,int(max_queries or 12))
    search_profile=build_search_profile(campaign)
    profile=Profile.objects.get_or_create(pk=1)[0]
    steering=search_profile.get('campaign_steering') or {}
    configured_roles=[_normalise(x) for x in (steering.get('roles') or []) if _normalise(x)]
    inferred_roles=[_normalise(x.get('role')) for x in (search_profile.get('role_families') or []) if _normalise(x.get('role'))]
    roles=configured_roles or inferred_roles[:4]
    engagement_terms=_engagement_context_terms(campaign,profile)
    llm=_local_query_language(campaign,profile,roles,engagement_terms)

    aliases=[]
    if not roles:
        aliases=[_normalise((x.get('term') if isinstance(x,dict) else x)) for x in (search_profile.get('skills') or [])[:6]]
        aliases=[x for x in aliases if x]
        if not aliases:
            return {'queries':[], 'search_profile':search_profile, 'location_policy':{'operating_location':profile.operating_location,'location_used_in_queries':False,'note':'Candidate Profile needs review before discovery can build profession-specific queries.'}, 'profile_review_required':True}
    for role in roles[:6]:
        role_aliases=[role]+list((llm.get('role_aliases') or {}).get(role,[]))+_role_alias_fallback(role)
        for alias in role_aliases:
            alias=_normalise(alias)
            if alias and alias.lower() not in {x.lower() for x in aliases}: aliases.append(alias)
    aliases=aliases[:24] or roles[:1]

    skills=search_profile.get('skills') or []

    def skill_term(row):
        return _normalise(row.get('term') if isinstance(row,dict) else str(row or ''))

    def rotate_plain(values, label):
        seen=[]
        for value in values:
            value=_normalise(value)
            if value and value.lower() not in {x.lower() for x in seen}:
                seen.append(value)
        rows=[{'query':f'{label}:{idx}:{value}','term':value} for idx,value in enumerate(seen)]
        salt=sum(ord(ch) for ch in str(label or ''))
        return [row['term'] for row in _stable_rotation(rows,getattr(campaign,'pk',None),interval_minutes,rotation_offset+salt)]

    cv_terms_all=[]; high_terms_all=[]; med_terms_all=[]; low_terms_all=[]; other_terms_all=[]
    for x in skills:
        term=skill_term(x)
        if not term:
            continue
        sources=x.get('sources',{}) if isinstance(x,dict) and isinstance(x.get('sources'),dict) else {}
        if sources.get('cv') or (isinstance(x,dict) and str(x.get('category') or '').casefold()=='cv-specific'):
            cv_terms_all.append(term)
        if sources.get('high_preference'):
            high_terms_all.append(term)
        elif sources.get('medium_preference'):
            med_terms_all.append(term)
        elif sources.get('low_preference'):
            low_terms_all.append(term)
        else:
            other_terms_all.append(term)
        if isinstance(x,dict):
            for variant in x.get('matched_variants') or []:
                variant=_normalise(variant)
                if variant and (sources.get('cv') or str(x.get('category') or '').casefold()=='cv-specific'):
                    cv_terms_all.append(variant)

    for x in _explicit_preference_phrases(profile.high_priority_text or ''):
        high_terms_all.append(x)
    for x in _explicit_preference_phrases(profile.medium_priority_text or ''):
        med_terms_all.append(x)
    for x in _explicit_preference_phrases(profile.low_priority_text or ''):
        low_terms_all.append(x)

    cv_terms=rotate_plain(cv_terms_all,'cv-tech')[:max(18,min(64,len(cv_terms_all) or 0))]
    high_terms=rotate_plain(high_terms_all,'high')[:18]
    med_terms=rotate_plain(med_terms_all,'medium')[:16]
    low_terms=rotate_plain(low_terms_all,'low')[:14]
    other_terms=rotate_plain(other_terms_all,'profile')[:18]
    tech=rotate_plain([_normalise(x) for x in (steering.get('technologies') or []) if _normalise(x)],'campaign-tech')[:16]
    # CV breadth is intentionally wide and rotated.  Interleave campaign and CV terms so
    # a long configured technology list cannot permanently block Resume/CV technologies
    # from reaching outgoing search queries.
    term_pool=[]
    grouped=[tech,cv_terms[:48],high_terms[:12],med_terms[:10],low_terms[:8],other_terms[:10]]
    for idx in range(max([len(x) for x in grouped] + [1])):
        for group in grouped:
            if idx < len(group) and group[idx].lower() not in {x.lower() for x in term_pool}:
                term_pool.append(group[idx])
    if not term_pool: term_pool=['specialist','technical']
    context=list(dict.fromkeys((llm.get('context_terms') or [])+engagement_terms+['hiring','looking for','startup','small team','early stage','new team','contractor','consultant','project']))[:24]

    candidates=[]
    def add(query,terms,kind,score,rationale):
        query=sanitize_local_search_engine_query(query)
        if not query: return
        loc=(search_profile.get('operating_location') or '').strip()
        if loc:
            query=re.sub(r'(?i)(?<![A-Za-z0-9])'+re.escape(loc)+r'(?![A-Za-z0-9])','',query)
            query=re.sub(r'\s+',' ',query).strip()
        candidates.append({'query':query,'terms':terms[:5],'kind':kind,'score':round(float(score),2),'rationale':rationale})

    remote_variants=['','', '', ' remote',' distributed',' "work from anywhere"']
    idx=0
    # Core role + specialist evidence. Pair terms where useful, but never produce an
    # unanchored technical query for a role-specific campaign.
    for ai,alias in enumerate(aliases):
        for ti,term in enumerate(term_pool[:max(24,min(48,len(term_pool)))]):
            second=''
            if (ti+ai)%3==0 and len(term_pool)>1:
                candidate=term_pool[(ti+5)%len(term_pool)]
                if candidate.lower()!=term.lower(): second=' '+_quote(candidate)
            remote=remote_variants[idx%len(remote_variants)]; idx+=1
            add(f'{_quote(alias)} {_quote(term)}{second}{remote}',[alias,term]+([second.strip(' \"')] if second else []),'role-specialist',20+max(0,12-ai/2-ti/5),f'Campaign role wording “{alias}” enriched with saved specialist evidence.')
            if len(candidates)>=max_queries*3: break
        if len(candidates)>=max_queries*3: break

    # Reserve explicit high/medium/low preference variants so lower tiers are not drowned out.
    tier_specs=[('high',high_terms,8),('medium',med_terms,5),('low',low_terms,2)]
    for tier,terms,boost in tier_specs:
        for ti,term in enumerate(terms[:8]):
            alias=aliases[(ti+len(tier))%len(aliases)]
            modifier=context[(ti+boost)%len(context)] if context else ''
            add(f'{_quote(alias)} {_quote(term)} {_quote(modifier)}'.strip(),[alias,term,modifier],f'priority-{tier}',18+boost-ti/4,f'{tier.title()}-priority preference injected while preserving the campaign role.')

    # Engagement/company-size settings and human-style opportunity language.
    for ci,modifier in enumerate(context[:18]):
        alias=aliases[ci%len(aliases)]; term=term_pool[(ci*3)%len(term_pool)]
        phrase='"looking for"' if ci%4==0 else ('hiring' if ci%4==1 else '')
        remote=remote_variants[(ci+2)%len(remote_variants)]
        add(f'{phrase} {_quote(alias)} {_quote(term)} {_quote(modifier)}{remote}'.strip(),[alias,term,modifier],'engagement-context',17-ci/10,'Campaign role + engagement/company-size/opportunity context.')

    # Target every configured custom domain when room permits, then rotate built-in source
    # domains. site: is explicit because many engines rank a source better when constrained.
    custom=[]
    for d in list(custom_domains or []):
        d=_normalise(str(d)).lower().removeprefix('www.')
        if d and d not in custom: custom.append(d)
    # Keep every enabled source identity in the scheduling pool. Domain-family expansion
    # happens only after source selection, eliminating the old 32-expanded-domain cutoff
    # that permanently excluded later sources such as YC Work at a Startup.
    source_targets=_source_target_rows(
        target_domains=target_domains,target_sources=target_sources,
        campaign_id=getattr(campaign,'pk',None),interval_minutes=interval_minutes,rotation_offset=rotation_offset,per_source_limit=6,
    )
    source_targets=[row for row in source_targets if row.get('source_domain') not in custom and row.get('domain') not in custom]
    # Every enabled Custom Domain gets its own explicit site: query before ordinary
    # source-domain rotation. This is stronger than broad search + post-filtering.
    custom_site_rows=[]
    for di,domain in enumerate(custom):
        alias=aliases[di%len(aliases)]; term=term_pool[di%len(term_pool)]
        custom_site_rows.append({'query':sanitize_local_search_engine_query(f'site:{domain} {_quote(alias)} {_quote(term)}'),'terms':[alias,term],'kind':'custom-domain','score':29-di/100,'rationale':f'Explicit Custom Domain search for {domain} anchored to campaign role wording.'})
    site_rows=[]
    for di,target in enumerate(source_targets):
        domain=target['domain']
        alias=aliases[di%len(aliases)]; term=term_pool[di%len(term_pool)]
        q=sanitize_local_search_engine_query(f'site:{domain} {_quote(alias)} {_quote(term)}')
        site_rows.append({
            'query':q,'terms':[alias,term],'kind':'site-targeted','score':25-di/100,
            'rationale':f'Campaign-role query explicitly targeted at configured source {target.get("source_name") or domain} via {domain}.',
            'target_source_id':target.get('source_id'),'target_source_name':target.get('source_name') or '',
            'target_source_domain':target.get('source_domain') or domain,'target_domain':domain,'source_surface':domain,
            'coverage_reason':'bounded-source-rotation',
        })

    unique={}
    for item in candidates:
        key=item['query'].lower()
        if key not in unique or item['score']>unique[key]['score']: unique[key]=item
    ranked=sorted(unique.values(),key=lambda x:-x['score'])
    # First reserve source-domain coverage (especially Custom Domains), then balanced priority
    # tiers, then fill from the rotating role-anchored pool.
    selected=[]
    # Reserve every Custom Domain first (up to the run's total query cap), then use a
    # bounded portion of the remaining slots for configured built-in source domains.
    for item in _stable_rotation(custom_site_rows,getattr(campaign,'pk',None),interval_minutes,rotation_offset):
        if len(selected)>=max_queries: break
        selected.append(item)
    remaining_slots=max(0,max_queries-len(selected))
    site_take=min(len(site_rows),max(1,round(remaining_slots*0.30))) if site_rows and remaining_slots else 0
    # Advance by the number of available source slots on each rotation. Unlike the old
    # hash-sort-after-truncation behavior, this gives a deterministic maximum starvation
    # bound of ceil(enabled_sources / site_take) rotations.
    for item in _coverage_rotation(site_rows,getattr(campaign,'pk',None),interval_minutes,site_take or 1,rotation_offset)[:site_take]:
        selected.append(item)
    for kind in ('priority-high','priority-medium','priority-low'):
        bucket=[x for x in ranked if x['kind']==kind]
        if bucket and len(selected)<max_queries:
            selected.append(_stable_rotation(bucket,getattr(campaign,'pk',None),interval_minutes,rotation_offset)[0])
    for item in _stable_rotation(ranked,getattr(campaign,'pk',None),interval_minutes,rotation_offset):
        if item not in selected: selected.append(item)
        if len(selected)>=max_queries: break
    if len(selected)<max_queries:
        for item in _coverage_rotation(site_rows,getattr(campaign,'pk',None),interval_minutes,max(1,len(site_rows)),rotation_offset):
            if item not in selected: selected.append(item)
            if len(selected)>=max_queries: break
    return {
        'queries':selected[:max_queries],
        'search_profile':search_profile,
        'candidate_query_count':len(ranked)+len(site_rows)+len(custom_site_rows),
        'location_used_in_queries':False,
        'location_policy':'Operating location is used for eligibility/fit after discovery, not blindly appended to search terms.',
        'query_policy':'Local AI Discovery role-anchored planning; at most three concise quoted expressions are sent per request; long exact phrases and Boolean filler such as AND/OR/NOT are removed; role aliases rotate; CV evidence plus preference tiers and engagement settings enrich queries; configured source domains use site: queries.',
        'rotation_offset':int(rotation_offset or 0),
        'rotation_interval_minutes':max(30,int(interval_minutes or 120)),
        'rotation_bucket':int(time.time()//(max(30,int(interval_minutes or 120))*60))+int(rotation_offset or 0),
        'role_aliases':aliases,
        'engagement_context':context,
        'custom_domains_targeted':custom,
        'source_target_count':len(site_rows),
        'source_site_slots':site_take,
        'source_coverage_max_rotations':((len(site_rows)+site_take-1)//site_take) if site_take else 0,
        'source_guided':True,
    }

def pre_score_hit(title: str, snippet: str, search_profile: dict) -> tuple[int, list[str]]:
    """Cheap pre-enrichment relevance score used during result consolidation."""
    blob = f'{title} {snippet}'.lower()
    score = 28.0
    matches = []
    for skill in search_profile.get('skills', [])[:PROFILE_CONCEPT_LIMIT]:
        terms = [skill['term']] + skill.get('aliases', []) + skill.get('matched_variants', [])
        if any(_has(blob, t) for t in terms):
            matches.append(skill['term'])
            score += min(8.0, 2.0 + skill.get('score', 0) / 6.0)
    if 'remote' in blob:
        score += 5
    if any(x in blob for x in ['contract', 'consultant', 'consulting']):
        score += 2
    for neg in search_profile.get('negative_terms', []):
        if neg.lower() in blob:
            score -= 10
    return max(0, min(100, int(round(score)))), list(dict.fromkeys(matches))[:10]


def grounded_profile_relevance(title: str, text: str, search_profile: dict, strict_market: bool=False) -> dict:
    """Validate discovery relevance against the fetched role page, not the search snippet.

    Search engines and multi-employer job boards often echo campaign terms in snippets,
    navigation and related-job widgets. Local AI Discovery therefore grounds relevance in
    the role title and the early role/JD body. On general job boards we use a stricter
    gate: an unrelated occupation title cannot be rescued by one specialist term appearing
    somewhere else on the page.
    """
    title_blob=_normalise(title or '')[:1200]
    full_body=_normalise(text or '')[:30000]
    # The first portion is much more likely to contain the actual role/JD than related jobs
    # and footer/navigation content appended later by aggregators.
    body_blob=full_body[:14000] if strict_market else full_body
    title_matches=[]; body_matches=[]; strong_body=[]
    weak_body_only={'python tooling','network protocols','c/c++','arm'}
    for skill in (search_profile or {}).get('skills',[])[:PROFILE_CONCEPT_LIMIT]:
        canonical=str(skill.get('term') or '').strip()
        if not canonical:
            continue
        terms=[canonical]+list(skill.get('aliases') or [])+list(skill.get('matched_variants') or [])
        title_hit=any(_has(title_blob,t) for t in terms if t)
        body_hit=title_hit or any(_has(body_blob,t) for t in terms if t)
        if title_hit:
            title_matches.append(canonical)
        if body_hit:
            body_matches.append(canonical)
            category=str(skill.get('category') or '')
            try: weight=float(skill.get('score') or 0)
            except Exception: weight=0
            if canonical.casefold() not in weak_body_only and (weight>=8 or category in {'reverse','virtualization','embedded-niche','retro-niche','protocol-niche','security','writing'}):
                strong_body.append(canonical)
    title_matches=list(dict.fromkeys(title_matches))
    body_matches=list(dict.fromkeys(body_matches))
    strong_body=list(dict.fromkeys(strong_body))
    technical_title=bool(re.search(r'(?i)\b(?:engineer|engineering|developer|programmer|architect|researcher|security|firmware|embedded|kernel|systems?|software|technical\s+(?:writer|author|trainer|instructor)|documentation\s+engineer|developer\s+documentation|devrel|reverse\s+engineer|virtuali[sz]ation)\b', title_blob))
    if strict_market:
        accepted=bool(title_matches or len(strong_body)>=2 or (technical_title and strong_body and len(body_matches)>=1))
    else:
        accepted=bool(title_matches or strong_body or len(body_matches)>=2)
    if title_matches:
        reason='Candidate-profile concept appears in the role title.'
    elif strict_market and len(strong_body)>=2:
        reason='Multiple strong candidate-profile concepts appear in the primary job description.'
    elif strict_market and technical_title and strong_body:
        reason='Technical role title plus grounded specialist evidence appears in the primary job description.'
    elif strong_body:
        reason='Strong candidate-profile concept appears in fetched role/JD content.'
    elif len(body_matches)>=2:
        reason='Multiple candidate-profile concepts appear in fetched role/JD content.'
    else:
        reason='No grounded candidate-profile relevance found in the fetched role title/JD; search-snippet and related-job overlap are ignored.'
    return {
        'accepted':accepted,
        'matches':body_matches[:12],
        'title_matches':title_matches[:8],
        'strong_matches':strong_body[:8],
        'technical_title':technical_title,
        'strict_market':bool(strict_market),
        'reason':reason,
    }
