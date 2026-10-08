import json
import html
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from urllib.parse import urlparse
from xml.etree import ElementTree as ET
from django.conf import settings
from django.utils import timezone
from portal.models import Application, Profile, PreparedApplicationFile
from .ai import generate
from .tracking import allocate_for_url, blog_base_url, resolve_article
from .cold import display_company_name
from .queryplanner import extract_active_cv_texts

EMAIL_RE=re.compile(r'[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}')


def _parse_generated_email(text, opportunity):
    lines=(text or '').strip().splitlines()
    subject=''
    if lines and lines[0].lower().startswith('subject:'):
        subject=lines[0].split(':',1)[1].strip()
        body='\n'.join(lines[1:]).strip()
    else:
        subject=f'{opportunity.title} — application'
        body=(text or '').strip()
    return subject[:500], body

def _recipient_name(opportunity):
    """Application email salutations use the company, not a potentially stale contact."""
    if not opportunity:
        return ''
    company=str(getattr(opportunity,'company','') or '').strip()
    url=str(getattr(opportunity,'target_url','') or getattr(opportunity,'canonical_url','') or getattr(opportunity,'url','') or '')
    return display_company_name(company,url) or company


def _cv_evidence(limit=14000):
    chunks=[]
    try:
        for doc in extract_active_cv_texts():
            text=' '.join(str(doc.get('text') or '').split())
            if text:
                chunks.append(f"Resume {doc.get('label') or doc.get('name') or ''}: {text}")
    except Exception:
        return ''
    return '\n'.join(chunks)[:limit]


def personalize_email_body(body, profile, opportunity=None):
    text=str(body or '')
    name=str(getattr(profile,'display_name','') or '').strip()
    phone=str(getattr(profile,'phone_number','') or '').strip()
    email=str(getattr(profile,'application_email','') or '').strip()
    recipient=_recipient_name(opportunity)
    if name:
        text=re.sub(r'\[\s*your\s+name\s*\]',lambda _m:name,text,flags=re.I)
        text=re.sub(r'\{\{\s*your(?:\s+|_)name\s*\}\}',lambda _m:name,text,flags=re.I)
        text=re.sub(r'\[\s*name\s*\]',lambda _m:name,text,flags=re.I)
    if phone:
        text=re.sub(r'\[\s*(?:contact|phone)\s+(?:number|no\.?|#)\s*\]',lambda _m:phone,text,flags=re.I)
        text=re.sub(r'\{\{\s*(?:contact|phone)(?:\s+|_)number\s*\}\}',lambda _m:phone,text,flags=re.I)
        text=re.sub(r'\[\s*your\s+contact\s+information\s*\]',lambda _m:phone,text,flags=re.I)
    elif email:
        text=re.sub(r'\[\s*your\s+contact\s+information\s*\]',lambda _m:email,text,flags=re.I)
    else:
        text=re.sub(r'(?im)^\s*\[\s*your\s+contact\s+information\s*\]\s*$', '', text)
    if email:
        text=re.sub(r'\[\s*(?:your\s+)?email(?:\s+address)?\s*\]',lambda _m:email,text,flags=re.I)
        text=re.sub(r'\{\{\s*(?:application(?:\s+|_)|your(?:\s+|_))?email\s*\}\}',lambda _m:email,text,flags=re.I)
    if recipient:
        text=re.sub(r"\[\s*recipient(?:'s|’s)?(?:\s*/\s*contact)?\s+name\s*\]",lambda _m:recipient,text,flags=re.I)
        text=re.sub(r'\[\s*recipient\s*/\s*contact\s+name\s*\]',lambda _m:recipient,text,flags=re.I)
        text=re.sub(r'\{\{\s*recipient(?:\s+|_)name\s*\}\}',lambda _m:recipient,text,flags=re.I)
    # Remove signature scaffolding that should never reach a real draft.
    text=re.sub(r'(?im)^\s*\[\s*your\s+position\s*\]\s*$', '', text)
    if name:
        text=re.sub(r'\[\s*'+re.escape(name)+r'\s*\]',lambda _m:name,text,flags=re.I)

    # Clean legacy/model leakage from older drafts. Instructional placeholders are
    # worse than an omitted sentence; real CV facts should come from the tailoring
    # prompt rather than prompting the user to fill examples manually.
    bad_placeholder=re.compile(r'\[(?:briefly\s+describe|describe\s+another|mention\s+specific|insert\b|example\b|your\s+(?:email|phone|name))[^]]*\]',re.I)
    cleaned=[]
    for line in text.splitlines():
        if bad_placeholder.search(line):
            continue
        cleaned.append(line)
    text='\n'.join(cleaned)
    text=re.sub(r'\n{3,}','\n\n',text).strip()
    return text


def _email_prompt(opportunity, profile, cv_evidence, alternate=False):
    company=_recipient_name(opportunity) or str(opportunity.company or '').strip() or 'Hiring Team'
    mode='Write a fresh alternate version of' if alternate else 'Write'
    return f'''{mode} a brief, concrete application/outreach email for the role below.

Hard requirements:
- Use ONLY facts present in JOB EVIDENCE, COMPANY PROFILE, CANDIDATE PROFILE, or RESUME EVIDENCE. Never invent an accomplishment, company initiative, product, tool, credential, eligibility claim, or number.
- Include 1-2 short, natural lines that refer to a concrete evidenced detail from COMPANY PROFILE (for example a real product, project, technical focus, market, or company activity), then connect it to specific relevant experience evidenced in the Candidate Profile/Resume. If COMPANY PROFILE has no concrete evidence, omit those lines entirely rather than making generic praise or guessing.
- Do not output any bracket placeholder, example placeholder, TODO, or generic bullet such as "[describe an accomplishment]". If evidence is unavailable, omit that sentence entirely.
- Greeting must be exactly "Dear {company},". Do not use a contact-person placeholder.
- Prefer 120-190 words. No "I hope this message finds you well", no generic excitement/filler, and no long skills inventory.
- Mention at most 2-3 concrete candidate facts that directly match this role, drawn from the Resume evidence.
- If a portfolio is useful, include the actual URL. If contact details are useful, use the actual phone/email below.
- End naturally with the candidate's real name.
- Return plain text beginning with one line: Subject: ...

ROLE: {opportunity.title}
COMPANY: {company}
JOB EVIDENCE:
{str(opportunity.description or '')[:7500]}

COMPANY PROFILE:
{json.dumps(opportunity.company_intel or {}, ensure_ascii=False, default=str)[:7000] or 'No collected company profile evidence is available. Do not invent company facts.'}

CANDIDATE PROFILE:
Name: {profile.display_name}
Phone: {profile.phone_number}
Application email: {profile.application_email}
Portfolio: {profile.portfolio_url}
High priority: {profile.high_priority_text}
Medium priority: {profile.medium_priority_text}
Low priority / avoid: {profile.low_priority_text}

RESUME EVIDENCE:
{cv_evidence or 'No active Resume text was available. Do not invent Resume facts.'}
'''


def prepare_application(opportunity, cv=None, cover=None, provider=None, model=None):
    app,_=Application.objects.get_or_create(opportunity=opportunity)
    if app.deleted_at:
        raise RuntimeError('This application/outreach record is in the Recycle Bin. Restore it before preparing it again.')
    if cv: app.cv=cv
    if cover: app.cover_letter=cover
    profile=Profile.objects.get_or_create(pk=1)[0]
    prompt=_email_prompt(opportunity,profile,_cv_evidence(),alternate=False)
    try:
        subject_type='outreach' if (opportunity.extracted_facts or {}).get('outreach') else 'application'
        response=generate(prompt,stage='email_draft',provider=provider,model=model,subject={'type':subject_type,'id':app.pk,'label':f'{opportunity.company} — {opportunity.title}'})
        subject,body=_parse_generated_email(response,opportunity)
        body=personalize_email_body(body,profile,opportunity)
    except Exception:
        company=_recipient_name(opportunity) or opportunity.company or 'Hiring Team'
        subject=f'{opportunity.title} — application'
        body=(f'Dear {company},\n\nI am writing about the {opportunity.title} role. '
              f'My background appears relevant to the work described, and I would welcome a conversation about where my experience could help.\n\n'
              + (f'Portfolio: {profile.portfolio_url}\n\n' if profile.portfolio_url else '')
              + f'Regards,\n{profile.display_name}')
    app.email_subject=subject; app.email_body=body; app.email_mode='plain'; app.status='prepared'; app.save()
    return app


def generate_email_version(application, provider=None, model=None):
    profile=Profile.objects.get_or_create(pk=1)[0]
    o=application.opportunity
    prompt=_email_prompt(o,profile,_cv_evidence(),alternate=True)
    subject_type='outreach' if (o.extracted_facts or {}).get('outreach') else 'application'
    text=generate(prompt,stage='email_draft',provider=provider,model=model,subject={'type':subject_type,'id':application.pk,'label':f'{o.company} — {o.title}'})
    subject,body=_parse_generated_email(text,o)
    return subject,personalize_email_body(body,profile,o)


def _replace_email_docx(src, dest, application_email):
    from docx import Document
    doc=Document(src)
    count=0
    def replace_runs(paragraphs):
        nonlocal count
        for p in paragraphs:
            for r in p.runs:
                found=EMAIL_RE.findall(r.text)
                if found:
                    count += len(found); r.text=EMAIL_RE.sub(application_email,r.text)
    replace_runs(doc.paragraphs)
    for section in doc.sections:
        replace_runs(section.header.paragraphs); replace_runs(section.footer.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                replace_runs(cell.paragraphs)
    doc.save(dest)
    return count


def _rewrite_docx_hyperlinks(docx_path, application):
    # Any hyperlink that resolves under the configured ToughDev blog base is handled
    # automatically. No manual article-title/vocabulary rule is required.
    tmp=Path(str(docx_path)+'.links.tmp')
    replacements=[]
    ns={'r':'http://schemas.openxmlformats.org/package/2006/relationships'}
    with zipfile.ZipFile(docx_path,'r') as zin, zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data=zin.read(item.filename)
            if item.filename.endswith('.rels') and item.filename.startswith('word/'):
                try:
                    root=ET.fromstring(data); changed=False
                    for rel in root.findall('r:Relationship',ns):
                        target=rel.attrib.get('Target','')
                        if rel.attrib.get('TargetMode')!='External' or not target: continue
                        try:
                            link,info=allocate_for_url(target,application)
                        except Exception:
                            continue
                        rel.set('Target',link.full_url)
                        replacements.append((target,link.full_url,info.get('title','')))
                        changed=True
                    if changed: data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
                except Exception:
                    pass
            zout.writestr(item,data)
    tmp.replace(docx_path)
    return replacements


def _to_pdf(docx_path):
    libreoffice=shutil.which('libreoffice') or shutil.which('soffice')
    if not libreoffice: return None,'LibreOffice not installed in image; PDF conversion skipped.'
    outdir=Path(docx_path).parent
    try:
        proc=subprocess.run([libreoffice,'--headless','--convert-to','pdf','--outdir',str(outdir),str(docx_path)],capture_output=True,text=True,timeout=60)
    except subprocess.TimeoutExpired:
        return None,'PDF conversion timed out after 60 seconds; the editable file is still available.'
    pdf=outdir/(Path(docx_path).stem+'.pdf')
    if proc.returncode==0 and pdf.exists(): return pdf,'PDF generated.'
    return None,('PDF conversion failed: '+(proc.stderr or proc.stdout)[-500:])


def generate_application_artifacts(application, kind='cv', progress=None, job_id=None):
    is_cover=(kind=='cover')
    asset=application.cover_letter if is_cover else application.cv
    label='Cover Letter' if is_cover else 'Resume'
    def update(value,message):
        if progress:
            progress(max(1,min(99,int(value))),message)
    if not asset: return {'ok':False,'message':f'No source {label} selected.'}
    update(8,f'Loading selected {label}')
    profile=Profile.objects.get_or_create(pk=1)[0]
    try:
        from .mailbox import active_profile
        ap=active_profile()
        application_email=(ap.imap_email if ap and ap.imap_email else profile.application_email)
    except Exception:
        application_email=profile.application_email
    src=Path(asset.file.path)
    stamp=timezone.now().strftime('%Y%m%d%H%M%S')
    outdir=Path(settings.MEDIA_ROOT)/'generated'/timezone.now().strftime('%Y/%m')
    outdir.mkdir(parents=True,exist_ok=True)
    stem='cover' if is_cover else 'resume'
    generated_field='generated_cover' if is_cover else 'generated_cv'
    generated_pdf_field='generated_cover_pdf' if is_cover else 'generated_cv_pdf'
    records=[]
    def remember(path,fmt,note=''):
        rel=str(Path(path).relative_to(settings.MEDIA_ROOT))
        rec=PreparedApplicationFile.objects.create(
            application=application,document_kind=kind,file_format=fmt,
            label=f'{label} {fmt.upper()} · {timezone.localtime().strftime("%d/%m/%Y %H:%M")}',
            file=rel,source_asset_label=asset.label,
            metadata={'note':note,'source_asset_id':asset.pk,'job_id':job_id},
        )
        records.append(rec.pk)
    if src.suffix.lower()!='.docx':
        update(45,f'Copying selected {label}')
        dest=outdir/f'app-{application.pk}-{stem}-{stamp}{src.suffix.lower()}'
        shutil.copy2(src,dest)
        setattr(application,generated_field,str(dest.relative_to(settings.MEDIA_ROOT)))
        application.save(update_fields=[generated_field,'updated_at'])
        remember(dest,src.suffix.lower().lstrip('.') or 'file','Source copied unchanged')
        update(95,f'{label} ready')
        return {'ok':True,'message':f'{label} source copied unchanged. Email/link rewriting requires a DOCX source.','path':str(dest),'links':[],'kind':kind,'prepared_file_ids':records}
    update(28,f'Updating {label} contact details')
    dest=outdir/f'app-{application.pk}-{stem}-{stamp}.docx'
    email_count=_replace_email_docx(src,dest,application_email)
    update(48,f'Updating {label} tracking links')
    links=_rewrite_docx_hyperlinks(dest,application)
    update(68,f'Creating {label} PDF')
    pdf,pdf_note=_to_pdf(dest)
    update(86,f'Recording prepared {label} files')
    setattr(application,generated_field,str(dest.relative_to(settings.MEDIA_ROOT)))
    if pdf: setattr(application,generated_pdf_field,str(pdf.relative_to(settings.MEDIA_ROOT)))
    application.save(update_fields=[generated_field,generated_pdf_field,'updated_at'])
    remember(dest,'docx',f'{email_count} email replacement(s); {len(links)} tracked link(s)')
    if pdf: remember(pdf,'pdf',pdf_note)
    update(96,f'{label} files ready')
    return {'ok':True,'message':f'{label} files ready; replaced {email_count} email occurrence(s); {len(links)} tracked link(s). {pdf_note}','path':str(dest),'pdf':str(pdf) if pdf else '', 'links':links,'kind':kind,'prepared_file_ids':records}


def generate_cv_artifacts(application):
    """Backward-compatible Resume artifact helper."""
    return generate_application_artifacts(application,'cv')


def scan_docx_links(file_obj, rules=None):
    """Return only hyperlinks that belong to the configured tracking blog.

    A DOCX can contain mailto links, reference sites and image relationships. The Tracking
    Link scanner is intentionally scoped to the configured Blog Base URL so unrelated links
    do not consume rows or appear as misleading "not trackable" results.
    """
    matches=[]
    base=urlparse(blog_base_url())
    base_host=(base.hostname or '').lower().removeprefix('www.')
    base_path=(base.path or '/').rstrip('/') or '/'
    with zipfile.ZipFile(file_obj,'r') as zin:
        ns={'r':'http://schemas.openxmlformats.org/package/2006/relationships'}
        for item in zin.infolist():
            if not (item.filename.endswith('.rels') and item.filename.startswith('word/')): continue
            try: root=ET.fromstring(zin.read(item.filename))
            except Exception: continue
            for rel in root.findall('r:Relationship',ns):
                target=rel.attrib.get('Target','')
                if rel.attrib.get('TargetMode')!='External': continue
                try:
                    parsed=urlparse(target)
                except Exception:
                    continue
                host=(parsed.hostname or '').lower().removeprefix('www.')
                path=(parsed.path or '/').rstrip('/') or '/'
                if not host or host!=base_host:
                    continue
                if base_path!='/' and not (path==base_path or path.startswith(base_path+'/')):
                    continue
                try:
                    info=resolve_article(target)
                    matches.append({'target':target,'would_track':True,'resolved_url':info.get('resolved_url',''),'article_title':info.get('title',''),'suffix_words':info.get('words',[])[:8]})
                except Exception as exc:
                    # Keep an on-blog broken/unresolvable URL visible so the Test action remains
                    # useful; only unrelated hosts/paths are suppressed.
                    matches.append({'target':target,'would_track':False,'resolved_url':'','article_title':'','suffix_words':[],'error':str(exc)})
    return matches


# Backward-compatible helper used by early view code.
def sync_cv_email(asset, application_email, destination):
    src=Path(asset.file.path); dest=Path(destination); dest.parent.mkdir(parents=True,exist_ok=True)
    if src.suffix.lower()!='.docx': shutil.copy2(src,dest); return 'PDF/non-DOCX copied unchanged'
    count=_replace_email_docx(src,dest,application_email)
    return f'Replaced {count} email occurrence(s)' if count else 'No email found; review generated Resume manually'
